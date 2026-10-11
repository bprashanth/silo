#!/usr/bin/env python3
"""Generate the deterministic field_log benchmark case."""

from __future__ import annotations

import json
import random
import shutil
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SEED = 731_905
BASE_DIR = Path(__file__).resolve().parent
FILES_DIR = BASE_DIR / "files"

LOCATIONS = {
    "Odisha": {
        "Koraput": ["Lamtaput", "Kakiriguma"],
        "Rayagada": ["Bissam Cuttack", "Muniguda"],
    },
    "Maharashtra": {
        "Gadchiroli": ["Bhilgaon", "Mendha"],
        "Nandurbar": ["Akkalkuwa", "Dhadgaon"],
    },
    "Karnataka": {
        "Raichur": ["Devadurga", "Manvi"],
        "Koppal": ["Gangavathi", "Kushtagi"],
    },
}
PROGRAMMES = [
    "Health & Nutrition",
    "Girls' Education",
    "Water Security",
    "Livelihoods",
]
FIELD_WORKERS = [
    "Anita Naik",
    "Farida Sheikh",
    "Kavita Jadhav",
    "Meena Majhi",
    "Prakash Hegde",
    "Ramesh Tekam",
    "Sushila Korram",
    "Yusuf Khan",
]


def build_clean_records() -> list[dict]:
    rng = random.Random(SEED)
    records: list[dict] = []
    start = date(2025, 4, 1)

    for index in range(480):
        # A real programme often has one especially active cluster; this also
        # guarantees that the village-ranking question has a unique winner.
        if index < 72:
            state, district, village = "Maharashtra", "Gadchiroli", "Bhilgaon"
        else:
            state = rng.choices(
                list(LOCATIONS), weights=[0.36, 0.34, 0.30], k=1
            )[0]
            district = rng.choice(list(LOCATIONS[state]))
            village = rng.choice(LOCATIONS[state][district])

        women = rng.randint(5, 38)
        men = rng.randint(2, 27)
        children = rng.randint(0, 31)
        records.append(
            {
                "visit_id": f"FL-2025-{index + 1:04d}",
                "visit_date": start + timedelta(days=rng.randint(0, 181)),
                "field_worker": rng.choice(FIELD_WORKERS),
                "state": state,
                "district": district,
                "village": village,
                "programme": rng.choices(
                    PROGRAMMES, weights=[0.31, 0.24, 0.22, 0.23], k=1
                )[0],
                "women": women,
                "men": men,
                "children": children,
                "follow_up": rng.random() < 0.29,
                "travel_inr": rng.randrange(450, 4_501, 50),
            }
        )
    return records


def derive_clean_answers(records: list[dict]) -> dict[str, int | float | str]:
    """Derive truth only from pristine records, before messy rows exist."""
    village_totals: defaultdict[str, int] = defaultdict(int)
    for record in records:
        village_totals[record["village"]] += (
            record["women"] + record["men"] + record["children"]
        )

    maximum = max(village_totals.values())
    winners = [name for name, value in village_totals.items() if value == maximum]
    if len(winners) != 1:
        raise RuntimeError(f"Village ranking is tied: {winners}")

    return {
        "q1": sum(
            record["women"] + record["men"] + record["children"]
            for record in records
        ),
        "q2": sum(
            record["women"]
            for record in records
            if record["state"] == "Odisha"
            and record["programme"] == "Health & Nutrition"
        ),
        "q3": sum(record["district"] == "Gadchiroli" for record in records),
        "q4": winners[0],
        "q5": round(
            100.0 * sum(record["follow_up"] for record in records) / len(records),
            2,
        ),
        "q6": sum(
            record["travel_inr"]
            for record in records
            if record["state"] == "Maharashtra"
            and record["programme"] == "Water Security"
        ),
    }


def messy_village(name: str, row_number: int) -> str:
    variants = (
        name.upper(),
        name.lower(),
        f"  {name} ",
        "   ".join(name.split()),
        name.swapcase(),
    )
    return variants[row_number % len(variants)]


def messy_travel(amount: int, row_number: int) -> int | str:
    if row_number % 5 == 0:
        return f"₹{amount:,}"
    if row_number % 7 == 0:
        return f" {amount:,} "
    return amount


def write_workbook(records: list[dict], output_path: Path) -> None:
    rng = random.Random(SEED + 1)
    duplicate_ids = {
        record["visit_id"] for record in rng.sample(records, k=58)
    }

    wb = Workbook()
    wb.properties.creator = "Sahyog Rural Initiatives"
    wb.properties.created = datetime(2025, 10, 1, 9, 0, 0)
    wb.properties.modified = datetime(2025, 10, 1, 9, 0, 0)
    ws = wb.active
    ws.title = "Field Visit Log"

    headers = [
        "Visit ID",
        "Visit Date",
        "Field Worker",
        "State",
        "District",
        "Village",
        "Programme",
        "Women Reached",
        "Men Reached",
        "Children Reached",
        "Follow-up Required",
        "Travel Reimbursement (INR)",
    ]
    ws.merge_cells("A1:L1")
    ws["A1"] = "Sahyog Rural Initiatives — Beneficiary Field Visit Log"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:L2")
    ws["A2"] = (
        "Data note: village spelling/case/spacing varies; duplicate Visit IDs "
        "are field resubmissions."
    )
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells("A3:L3")
    ws["A3"] = "Reporting period: April–September 2025 | Amounts are INR"

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    output_row = 6
    emitted_data_rows = 0
    for record in records:
        copies = 2 if record["visit_id"] in duplicate_ids else 1
        for copy_index in range(copies):
            emitted_data_rows += 1
            values = [
                record["visit_id"],
                record["visit_date"],
                record["field_worker"],
                record["state"],
                record["district"],
                messy_village(
                    record["village"], emitted_data_rows + 2 * copy_index
                ),
                record["programme"],
                record["women"],
                record["men"],
                record["children"],
                " YES " if record["follow_up"] and emitted_data_rows % 4 == 0
                else "Yes" if record["follow_up"]
                else " no " if emitted_data_rows % 6 == 0
                else "No",
                messy_travel(record["travel_inr"], emitted_data_rows),
            ]
            for column, value in enumerate(values, start=1):
                ws.cell(row=output_row, column=column, value=value)
            ws.cell(row=output_row, column=2).number_format = "dd-mmm-yyyy"
            output_row += 1

        if emitted_data_rows % 93 == 0:
            output_row += 1

    widths = [17, 14, 19, 14, 15, 21, 21, 15, 13, 17, 20, 26]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:L{output_row - 1}"
    wb.save(output_path)


def build_questions(answers: dict[str, int | float | str]) -> list[dict]:
    common = (
        "Normalise village names by trimming outer whitespace, collapsing "
        "internal whitespace and comparing case-insensitively; deduplicate "
        "Visit ID by keeping its first spreadsheet occurrence. "
    )
    return [
        {
            "id": "q1",
            "question": common
            + "What is the total number of people reached (women + men + children)?",
            "answer": answers["q1"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Counting duplicate Visit IDs inflates the total reached.",
        },
        {
            "id": "q2",
            "question": common
            + "How many women were reached by Health & Nutrition visits in Odisha?",
            "answer": answers["q2"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Summing all Odisha attendees or retaining resubmitted visits is wrong.",
        },
        {
            "id": "q3",
            "question": common
            + "How many distinct retained visits took place in Gadchiroli district?",
            "answer": answers["q3"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Spreadsheet row count includes duplicate Visit IDs and blank separators.",
        },
        {
            "id": "q4",
            "question": common
            + "Which village reached the most people in total? Report the normalised village in title case.",
            "answer": answers["q4"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "Case and spacing variants otherwise split one village into several groups.",
        },
        {
            "id": "q5",
            "question": common
            + "What percentage of retained visits require follow-up? Round to two decimal places.",
            "answer": answers["q5"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Using raw rows as the denominator gives duplicates extra weight.",
        },
        {
            "id": "q6",
            "question": common
            + "Treat numeric and ₹/comma-formatted travel cells as INR. What was the total travel reimbursement for Water Security visits in Maharashtra?",
            "answer": answers["q6"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Text-formatted INR cells are easily dropped, and duplicate visits are easily counted twice.",
        },
    ]


def main() -> None:
    clean_records = build_clean_records()
    answers = derive_clean_answers(clean_records)

    if FILES_DIR.exists():
        shutil.rmtree(FILES_DIR)
    FILES_DIR.mkdir(parents=True)
    workbook_path = FILES_DIR / "field_visit_log.xlsx"
    write_workbook(clean_records, workbook_path)

    case = {
        "id": "field_log",
        "description": (
            "A field-staff beneficiary visit log with inconsistent village "
            "case/spacing, duplicate Visit IDs, blank separators and mixed INR cells."
        ),
        "files": [workbook_path.name],
        "questions": build_questions(answers),
    }
    (BASE_DIR / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    if workbook_path.stat().st_size >= 2 * 1024 * 1024:
        raise RuntimeError(f"{workbook_path.name} exceeds 2 MB")
    print(
        f"field_log: generated {workbook_path.name} "
        f"({workbook_path.stat().st_size} bytes), {len(case['questions'])} questions"
    )


if __name__ == "__main__":
    main()
