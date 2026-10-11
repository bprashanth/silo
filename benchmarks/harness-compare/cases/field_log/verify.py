#!/usr/bin/env python3
"""Independently verify every field_log answer from the messy workbook."""

from __future__ import annotations

import json
import math
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
EXPECTED_IDS = {"q1", "q2", "q3", "q4", "q5", "q6"}


def derive_with_openpyxl(path: Path) -> dict[str, int | float | str]:
    """First derivation: cell-by-cell parsing and ordinary Python reductions."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Field Visit Log"]

    header_row = None
    headers: dict[str, int] = {}
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        labels = [str(value).strip() if value is not None else "" for value in row]
        if "Visit ID" in labels:
            header_row = row_number
            headers = {label: index for index, label in enumerate(labels) if label}
            break
    if header_row is None:
        raise ValueError("Could not locate Visit ID header")

    required = {
        "Visit ID",
        "State",
        "District",
        "Village",
        "Programme",
        "Women Reached",
        "Men Reached",
        "Children Reached",
        "Follow-up Required",
        "Travel Reimbursement (INR)",
    }
    missing = required - headers.keys()
    if missing:
        raise ValueError(f"Missing workbook columns: {sorted(missing)}")

    retained: dict[str, dict] = {}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_id = row[headers["Visit ID"]]
        visit_id = "" if raw_id is None else str(raw_id).strip()
        if not visit_id.startswith("FL-"):
            continue
        if visit_id in retained:
            continue

        def integer(column: str) -> int:
            value = row[headers[column]]
            if isinstance(value, bool):
                raise ValueError(f"Boolean in numeric field {column}: {visit_id}")
            return int(value)

        raw_village = str(row[headers["Village"]])
        village_key = " ".join(raw_village.strip().split()).casefold()
        travel_value = row[headers["Travel Reimbursement (INR)"]]
        if isinstance(travel_value, (int, float)) and not isinstance(
            travel_value, bool
        ):
            travel = int(travel_value)
        else:
            cleaned = re.sub(r"[₹,\s]", "", str(travel_value))
            travel = int(cleaned)

        retained[visit_id] = {
            "state": str(row[headers["State"]]).strip(),
            "district": str(row[headers["District"]]).strip(),
            "village": village_key,
            "programme": str(row[headers["Programme"]]).strip(),
            "women": integer("Women Reached"),
            "men": integer("Men Reached"),
            "children": integer("Children Reached"),
            "follow_up": str(row[headers["Follow-up Required"]])
            .strip()
            .casefold()
            == "yes",
            "travel": travel,
        }
    workbook.close()

    if not retained:
        raise ValueError("No field visit records parsed")

    village_totals: defaultdict[str, int] = defaultdict(int)
    for record in retained.values():
        village_totals[record["village"]] += (
            record["women"] + record["men"] + record["children"]
        )
    maximum = max(village_totals.values())
    winners = [name for name, amount in village_totals.items() if amount == maximum]
    if len(winners) != 1:
        raise ValueError(f"Village ranking is not unique: {winners}")

    return {
        "q1": sum(
            record["women"] + record["men"] + record["children"]
            for record in retained.values()
        ),
        "q2": sum(
            record["women"]
            for record in retained.values()
            if record["state"] == "Odisha"
            and record["programme"] == "Health & Nutrition"
        ),
        "q3": sum(
            record["district"] == "Gadchiroli" for record in retained.values()
        ),
        "q4": winners[0].title(),
        "q5": round(
            100.0
            * sum(record["follow_up"] for record in retained.values())
            / len(retained),
            2,
        ),
        "q6": sum(
            record["travel"]
            for record in retained.values()
            if record["state"] == "Maharashtra"
            and record["programme"] == "Water Security"
        ),
    }


def raw_xlsx_rows(path: Path) -> list[dict[int, object]]:
    """Parse the first worksheet directly from XLSX XML, without openpyxl."""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    qualified = lambda name: f"{{{namespace}}}{name}"
    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(qualified("si")):
                shared_strings.append(
                    "".join(
                        node.text or "" for node in item.iter(qualified("t"))
                    )
                )
        sheet_root = ElementTree.fromstring(
            archive.read("xl/worksheets/sheet1.xml")
        )

    parsed_rows: list[dict[int, object]] = []
    for xml_row in sheet_root.findall(
        f".//{qualified('sheetData')}/{qualified('row')}"
    ):
        parsed: dict[int, object] = {}
        for cell in xml_row.findall(qualified("c")):
            reference = cell.attrib["r"]
            letters = re.match(r"[A-Z]+", reference)
            if letters is None:
                raise ValueError(f"Bad XLSX cell reference: {reference}")
            column = 0
            for letter in letters.group(0):
                column = column * 26 + ord(letter) - ord("A") + 1

            cell_type = cell.attrib.get("t", "n")
            if cell_type == "inlineStr":
                value: object = "".join(
                    node.text or "" for node in cell.iter(qualified("t"))
                )
            else:
                value_node = cell.find(qualified("v"))
                raw = "" if value_node is None else value_node.text or ""
                if cell_type == "s":
                    value = shared_strings[int(raw)]
                elif cell_type in {"str", "e"}:
                    value = raw
                elif cell_type == "b":
                    value = raw == "1"
                elif raw == "":
                    value = None
                else:
                    number = float(raw)
                    value = int(number) if number.is_integer() else number
            parsed[column] = value
        parsed_rows.append(parsed)
    return parsed_rows


def derive_with_xml(path: Path) -> dict[str, int | float | str]:
    """Second derivation: raw XLSX XML parsing and streaming accumulators."""
    rows = raw_xlsx_rows(path)
    header_position = None
    columns: dict[str, int] = {}
    for position, row in enumerate(rows):
        labels = {
            str(value).strip(): column
            for column, value in row.items()
            if value is not None
        }
        if "Visit ID" in labels:
            header_position = position
            columns = labels
            break
    if header_position is None:
        raise ValueError("XML parser could not locate Visit ID header")

    seen: set[str] = set()
    village_totals: defaultdict[str, int] = defaultdict(int)
    total_people = odisha_health_women = gadchiroli_visits = 0
    follow_up_visits = maharashtra_water_travel = 0

    for row in rows[header_position + 1 :]:
        visit_id = str(row.get(columns["Visit ID"], "")).strip()
        if not visit_id.startswith("FL-") or visit_id in seen:
            continue
        seen.add(visit_id)

        state = str(row[columns["State"]]).strip()
        district = str(row[columns["District"]]).strip()
        programme = str(row[columns["Programme"]]).strip()
        village = " ".join(str(row[columns["Village"]]).strip().split()).casefold()
        women = int(row[columns["Women Reached"]])
        people = (
            women
            + int(row[columns["Men Reached"]])
            + int(row[columns["Children Reached"]])
        )
        raw_travel = row[columns["Travel Reimbursement (INR)"]]
        travel = (
            int(raw_travel)
            if isinstance(raw_travel, (int, float)) and not isinstance(raw_travel, bool)
            else int(re.sub(r"[₹,\s]", "", str(raw_travel)))
        )

        total_people += people
        village_totals[village] += people
        if state == "Odisha" and programme == "Health & Nutrition":
            odisha_health_women += women
        if district == "Gadchiroli":
            gadchiroli_visits += 1
        if (
            str(row[columns["Follow-up Required"]]).strip().casefold()
            == "yes"
        ):
            follow_up_visits += 1
        if state == "Maharashtra" and programme == "Water Security":
            maharashtra_water_travel += travel

    maximum = max(village_totals.values())
    winners = [
        village for village, amount in village_totals.items() if amount == maximum
    ]
    if len(winners) != 1:
        raise ValueError(f"Village ranking is not unique: {winners}")
    return {
        "q1": total_people,
        "q2": odisha_health_women,
        "q3": gadchiroli_visits,
        "q4": winners[0].title(),
        "q5": round(100.0 * follow_up_visits / len(seen), 2),
        "q6": maharashtra_water_travel,
    }


def answers_match(actual: object, expected: object, kind: str, tol_pct: float) -> bool:
    if kind == "name":
        return isinstance(actual, str) and actual == expected
    if not isinstance(actual, (int, float)) or isinstance(actual, bool):
        return False
    if not isinstance(expected, (int, float)) or isinstance(expected, bool):
        return False
    allowed = abs(float(expected)) * tol_pct / 100.0
    return math.isclose(
        float(actual), float(expected), rel_tol=0.0, abs_tol=max(allowed, 1e-9)
    )


def main() -> int:
    try:
        case = json.loads((BASE_DIR / "case.json").read_text(encoding="utf-8"))
        questions = case.get("questions", [])
        if len(questions) not in (5, 6):
            raise ValueError(f"Expected 5 or 6 questions, found {len(questions)}")
        ids = {question.get("id") for question in questions}
        if ids != EXPECTED_IDS:
            raise ValueError(f"Unexpected question IDs: {sorted(ids)}")
        if case.get("files") != ["field_visit_log.xlsx"]:
            raise ValueError(f"Unexpected files declaration: {case.get('files')}")

        workbook_path = BASE_DIR / "files" / "field_visit_log.xlsx"
        if workbook_path.stat().st_size >= 2 * 1024 * 1024:
            raise ValueError("Generated workbook is 2 MB or larger")

        first = derive_with_openpyxl(workbook_path)
        second = derive_with_xml(workbook_path)
        errors: list[str] = []
        for question in questions:
            question_id = question["id"]
            if not answers_match(
                first[question_id],
                second[question_id],
                question["kind"],
                0.0,
            ):
                errors.append(
                    f"{question_id}: independent derivations disagree "
                    f"({first[question_id]!r} != {second[question_id]!r})"
                )
            if not answers_match(
                first[question_id],
                question["answer"],
                question["kind"],
                float(question.get("tol_pct", 0.0)),
            ):
                errors.append(
                    f"{question_id}: case answer {question['answer']!r}, "
                    f"parsed answer {first[question_id]!r}"
                )

        if errors:
            for error in errors:
                print(f"ERROR: {error}", file=sys.stderr)
            return 1
        print(
            f"field_log: PASS — {len(questions)} questions verified by "
            "independent openpyxl and raw-XML derivations"
        )
        return 0
    except Exception as exc:
        print(f"field_log: FAIL — {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
