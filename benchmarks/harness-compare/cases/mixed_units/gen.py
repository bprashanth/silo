#!/usr/bin/env python3
"""Generate the deterministic mixed_units benchmark case."""

from __future__ import annotations

import json
import random
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SEED = 20250403
ROOT = Path(__file__).resolve().parent
FILES = ROOT / "files"
WORKBOOK_NAME = "mixed_units_register.xlsx"

PROGRAMMES = [
    "Maternal & Child Health",
    "Rural Livelihoods",
    "Tribal Education",
    "Water & Sanitation",
]
STATES = ["Odisha", "Karnataka", "Rajasthan", "Jharkhand"]
DONORS = [
    "Azim Premji Philanthropic Initiatives",
    "Tata CSR Trust",
    "National Health Mission",
    "Global Community Foundation",
]
PURPOSES = [
    "Community mobilisation",
    "Training and workshops",
    "Equipment procurement",
    "Field team support",
    "Direct programme grant",
]


def make_clean_records() -> list[dict]:
    rng = random.Random(SEED)
    start = date(2025, 4, 1)
    records = []
    variants = [
        ("number", None),
        ("rupee", 120_000),
        ("short_l", 120_000),
        ("lakh", 250_000),
        ("crore", 10_000_000),
        ("k", 45_000),
        ("blank", 0),
        ("dash", 0),
        ("na", 0),
    ]
    numeric_values = [32_500, 78_000, 165_000, 340_000, 875_000]

    for index in range(180):
        variant, fixed_amount = variants[index % len(variants)]
        amount = rng.choice(numeric_values) if fixed_amount is None else fixed_amount
        record_date = (
            date(2025, 4, 3)
            if index < 3
            else start + timedelta(days=rng.randrange(365))
        )
        records.append(
            {
                "record_id": f"MU-{index + 1:04d}",
                "date": record_date,
                "programme": rng.choice(PROGRAMMES),
                "state": rng.choice(STATES),
                "donor": rng.choice(DONORS),
                "amount": amount,
                "amount_variant": variant,
                "date_variant": index % 3,
                "purpose": rng.choice(PURPOSES),
            }
        )
    return records


def compute_truth(records: list[dict]) -> dict[str, int | float | str]:
    total = sum(row["amount"] for row in records)
    q1_end = date(2025, 6, 30)
    q2 = sum(
        row["amount"]
        for row in records
        if row["programme"] == "Maternal & Child Health"
        and date(2025, 4, 1) <= row["date"] <= q1_end
    )
    q3 = sum(
        1 for row in records if row["state"] == "Odisha" and row["amount"] > 0
    )

    donor_totals = defaultdict(int)
    programme_totals = defaultdict(int)
    for row in records:
        donor_totals[row["donor"]] += row["amount"]
        programme_totals[row["programme"]] += row["amount"]
    q4 = max(donor_totals, key=donor_totals.get)
    q5 = 100.0 * sum(
        row["amount"] for row in records if row["state"] == "Karnataka"
    ) / total
    q6 = max(programme_totals, key=programme_totals.get)

    assert list(donor_totals.values()).count(donor_totals[q4]) == 1
    assert list(programme_totals.values()).count(programme_totals[q6]) == 1
    return {
        "q1": total,
        "q2": q2,
        "q3": q3,
        "q4": q4,
        "q5": round(q5, 6),
        "q6": q6,
    }


def display_amount(row: dict) -> int | str | None:
    variant = row["amount_variant"]
    values = {
        "rupee": "₹1,20,000",
        "short_l": "1.2 L",
        "lakh": "2.5 lakh",
        "crore": "1 Cr",
        "k": "45k",
        "blank": None,
        "dash": "-",
        "na": "NA",
    }
    return row["amount"] if variant == "number" else values[variant]


def display_date(row: dict) -> str:
    value = row["date"]
    if row["date_variant"] == 0:
        return value.strftime("%Y-%m-%d")
    if row["date_variant"] == 1:
        return value.strftime("%d/%m/%Y")
    return f"{value.day} {value.strftime('%b %Y')}"


def write_workbook(records: list[dict]) -> None:
    FILES.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Disbursement Register"
    ws.merge_cells("A1:G1")
    ws["A1"] = "Sahyog Foundation — FY 2025-26 Disbursement Register"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Amount note:"
    ws["B2"] = (
        "Numbers are INR. L and lakh mean 100,000; Cr means 10,000,000; "
        "k means 1,000. Blank, '-' and NA are zero."
    )
    ws.merge_cells("B2:G2")
    ws["A3"] = "Date note:"
    ws["B3"] = (
        "Dates use YYYY-MM-DD, DD/MM/YYYY (day first), or D Mon YYYY."
    )
    ws.merge_cells("B3:G3")
    ws["A4"] = "Prepared for:"
    ws["B4"] = "Programme management review — 31 March 2026"

    headers = [
        "Record ID",
        "Disbursement Date",
        "Programme",
        "State",
        "Donor / Scheme",
        "Amount",
        "Purpose",
    ]
    ws.append([])
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[6]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for index, row in enumerate(records, start=1):
        if index > 1 and (index - 1) % 45 == 0:
            ws.append([])
        if index > 1 and (index - 1) % 60 == 0:
            ws.append(headers)
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
        ws.append(
            [
                row["record_id"],
                display_date(row),
                row["programme"],
                row["state"],
                row["donor"],
                display_amount(row),
                row["purpose"],
            ]
        )

    widths = [14, 20, 28, 15, 40, 18, 28]
    for column, width in zip("ABCDEFG", widths):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:G{ws.max_row}"
    wb.save(FILES / WORKBOOK_NAME)


def make_case(truth: dict[str, int | float | str]) -> dict:
    amount_rule = (
        "Use INR; L/lakh = 100,000, Cr = 10,000,000, k = 1,000, "
        "and treat blank, '-' and NA amounts as zero."
    )
    return {
        "id": "mixed_units",
        "description": (
            "A disbursement register with mixed INR units, three date formats, "
            "blank separators, and repeated headers."
        ),
        "files": [WORKBOOK_NAME],
        "questions": [
            {
                "id": "q1",
                "question": f"What is the total disbursed amount? {amount_rule}",
                "answer": truth["q1"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Treating lakh/crore suffixes as literal numbers or dropping formatted strings.",
            },
            {
                "id": "q2",
                "question": (
                    "How much was disbursed to Maternal & Child Health from "
                    "1 April through 30 June 2025 inclusive? Dates are YYYY-MM-DD, "
                    "DD/MM/YYYY (day first), or D Mon YYYY. " + amount_rule
                ),
                "answer": truth["q2"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Parsing 03/04/2025 as March 4 instead of 3 April.",
            },
            {
                "id": "q3",
                "question": (
                    "How many Odisha records have a non-zero amount? "
                    "Treat blank, '-' and NA as zero and count records, not distinct amounts."
                ),
                "answer": truth["q3"],
                "kind": "number",
                "tol_pct": 0,
                "trap": "Counting missing-value markers as non-zero text entries.",
            },
            {
                "id": "q4",
                "question": (
                    "Which donor or scheme has the highest total disbursed amount? "
                    + amount_rule
                ),
                "answer": truth["q4"],
                "kind": "name",
                "tol_pct": 0,
                "trap": "Comparing formatted amount strings lexicographically.",
            },
            {
                "id": "q5",
                "question": (
                    "What percentage of the total disbursed amount went to Karnataka? "
                    "Return a percentage from 0 to 100. " + amount_rule
                ),
                "answer": truth["q5"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Returning a fraction instead of a percentage or omitting crore rows.",
            },
            {
                "id": "q6",
                "question": (
                    "Which programme received the highest total disbursed amount? "
                    + amount_rule
                ),
                "answer": truth["q6"],
                "kind": "name",
                "tol_pct": 0,
                "trap": "Including repeated header rows as data or failing to normalize units.",
            },
        ],
    }


def main() -> None:
    records = make_clean_records()
    truth = compute_truth(records)
    case = make_case(truth)
    write_workbook(records)
    (ROOT / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        f"mixed_units: generated {len(records)} records, "
        f"{len(case['questions'])} questions"
    )


if __name__ == "__main__":
    main()
