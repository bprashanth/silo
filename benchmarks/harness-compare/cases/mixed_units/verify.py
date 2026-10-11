#!/usr/bin/env python3
"""Independently verify every mixed_units answer from the messy workbook."""

from __future__ import annotations

import json
import math
import posixpath
import re
import sys
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "files" / "mixed_units_register.xlsx"
ID_PATTERN = re.compile(r"^MU-\d{4}$")


def amount_openpyxl(value) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    text = str(value).strip()
    if not text or text.upper() == "NA" or text == "-":
        return 0
    compact = text.replace("₹", "").replace(",", "").strip().lower()
    suffixes = [
        ("lakh", 100_000),
        ("cr", 10_000_000),
        ("l", 100_000),
        ("k", 1_000),
    ]
    for suffix, multiplier in suffixes:
        if compact.endswith(suffix):
            number = compact[: -len(suffix)].strip()
            return int(round(float(number) * multiplier))
    return int(round(float(compact)))


def date_openpyxl(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    formats = (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d %b %Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"unrecognised date: {value!r}")


def derive_with_openpyxl() -> dict[str, int | float | str]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    ws = wb["Disbursement Register"]
    records = []
    for values in ws.iter_rows(values_only=True):
        record_id = str(values[0]).strip() if values and values[0] is not None else ""
        if not ID_PATTERN.fullmatch(record_id):
            continue
        records.append(
            {
                "date": date_openpyxl(values[1]),
                "programme": str(values[2]).strip(),
                "state": str(values[3]).strip(),
                "donor": str(values[4]).strip(),
                "amount": amount_openpyxl(values[5]),
            }
        )
    wb.close()

    total = sum(item["amount"] for item in records)
    april_to_june = sum(
        item["amount"]
        for item in records
        if item["programme"] == "Maternal & Child Health"
        and date(2025, 4, 1) <= item["date"] <= date(2025, 6, 30)
    )
    odisha_nonzero = sum(
        item["state"] == "Odisha" and item["amount"] > 0 for item in records
    )
    donor_totals = defaultdict(int)
    programme_totals = defaultdict(int)
    for item in records:
        donor_totals[item["donor"]] += item["amount"]
        programme_totals[item["programme"]] += item["amount"]
    karnataka = sum(
        item["amount"] for item in records if item["state"] == "Karnataka"
    )
    return {
        "q1": total,
        "q2": april_to_june,
        "q3": odisha_nonzero,
        "q4": max(donor_totals.items(), key=lambda pair: pair[1])[0],
        "q5": round(100.0 * karnataka / total, 6),
        "q6": max(programme_totals.items(), key=lambda pair: pair[1])[0],
    }


def amount_ooxml(value) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(round(float(value)))
    token = str(value).strip().casefold()
    if token in {"", "-", "na"}:
        return 0
    token = token.replace("₹", "").replace(",", "").replace(" ", "")
    match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)(lakh|cr|l|k)?", token)
    if not match:
        raise ValueError(f"unrecognised amount: {value!r}")
    multipliers = {None: 1, "lakh": 100_000, "l": 100_000, "cr": 10_000_000, "k": 1_000}
    return int(round(float(match.group(1)) * multipliers[match.group(2)]))


def date_ooxml(value) -> date:
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return datetime.strptime(text, "%Y-%m-%d").date()
    if "/" in text:
        return datetime.strptime(text, "%d/%m/%Y").date()
    return datetime.strptime(text, "%d %b %Y").date()


def column_index(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference).group(0)
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result - 1


def read_ooxml_rows(sheet_name: str) -> list[list[object]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(WORKBOOK) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relation_id = None
        for sheet in workbook_root.findall(f".//{{{main_ns}}}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                relation_id = sheet.attrib[f"{{{office_rel_ns}}}id"]
                break
        if relation_id is None:
            raise ValueError(f"sheet not found: {sheet_name}")
        relations = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = None
        for relation in relations.findall(f"{{{package_rel_ns}}}Relationship"):
            if relation.attrib.get("Id") == relation_id:
                target = relation.attrib["Target"]
                break
        if target is None:
            raise ValueError(f"relationship not found for sheet: {sheet_name}")
        sheet_path = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))
        sheet_root = ElementTree.fromstring(archive.read(sheet_path))
        output = []
        for xml_row in sheet_root.findall(f".//{{{main_ns}}}row"):
            values: list[object] = [None] * 7
            for cell in xml_row.findall(f"{{{main_ns}}}c"):
                index = column_index(cell.attrib["r"])
                if index >= len(values):
                    continue
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    inline = cell.find(f"{{{main_ns}}}is")
                    value = "".join(
                        node.text or "" for node in inline.iter(f"{{{main_ns}}}t")
                    )
                else:
                    value_node = cell.find(f"{{{main_ns}}}v")
                    if value_node is None:
                        value = None
                    elif cell_type == "s":
                        value = shared[int(value_node.text)]
                    elif cell_type in ("str", "e"):
                        value = value_node.text
                    else:
                        number = float(value_node.text)
                        value = int(number) if number.is_integer() else number
                values[index] = value
            output.append(values)
        return output


def derive_with_ooxml() -> dict[str, int | float | str]:
    rows = []
    for row in read_ooxml_rows("Disbursement Register"):
        key = "" if row[0] is None else str(row[0]).strip()
        if not ID_PATTERN.fullmatch(key):
            continue
        rows.append(
            (
                date_ooxml(row[1]),
                str(row[2]).strip(),
                str(row[3]).strip(),
                str(row[4]).strip(),
                amount_ooxml(row[5]),
            )
        )

    total = sum(row[4] for row in rows)
    q2 = sum(
        row[4]
        for row in rows
        if row[1] == "Maternal & Child Health"
        and date(2025, 4, 1) <= row[0] <= date(2025, 6, 30)
    )
    q3 = len([row for row in rows if row[2] == "Odisha" and row[4] != 0])
    by_donor: dict[str, int] = {}
    by_programme: dict[str, int] = {}
    for _, programme, _, donor, amount in rows:
        by_donor[donor] = by_donor.get(donor, 0) + amount
        by_programme[programme] = by_programme.get(programme, 0) + amount
    q4 = sorted(by_donor, key=lambda key: (-by_donor[key], key))[0]
    q5 = round(
        100.0 * sum(row[4] for row in rows if row[2] == "Karnataka") / total, 6
    )
    q6 = sorted(by_programme, key=lambda key: (-by_programme[key], key))[0]
    return {"q1": total, "q2": q2, "q3": q3, "q4": q4, "q5": q5, "q6": q6}


def values_match(left, right, tol_pct: float = 0.0) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        return left == right
    left_number = float(left)
    right_number = float(right)
    if tol_pct == 0:
        return math.isclose(left_number, right_number, rel_tol=0, abs_tol=1e-9)
    scale = max(abs(right_number), 1.0)
    return abs(left_number - right_number) <= scale * tol_pct / 100.0


def main() -> int:
    try:
        case = json.loads((ROOT / "case.json").read_text(encoding="utf-8"))
        first = derive_with_openpyxl()
        second = derive_with_ooxml()
        errors = []

        if set(first) != set(second):
            errors.append("independent derivations produced different question sets")
        for question_id in sorted(set(first) | set(second)):
            if question_id not in first or question_id not in second:
                continue
            if not values_match(first[question_id], second[question_id]):
                errors.append(
                    f"{question_id}: derivations disagree "
                    f"({first[question_id]!r} vs {second[question_id]!r})"
                )

        questions = case.get("questions", [])
        if len(questions) not in (5, 6):
            errors.append(f"case has {len(questions)} questions; expected 5 or 6")
        for question in questions:
            question_id = question["id"]
            if question_id not in first:
                errors.append(f"{question_id}: no derived answer")
                continue
            if not values_match(
                first[question_id], question["answer"], float(question.get("tol_pct", 0))
            ):
                errors.append(
                    f"{question_id}: derived {first[question_id]!r}, "
                    f"case.json has {question['answer']!r}"
                )

        if errors:
            for error in errors:
                print(f"ERROR: {error}", file=sys.stderr)
            return 1
        print(
            f"mixed_units verify: PASS ({len(questions)} questions; "
            "openpyxl and OOXML derivations agree)"
        )
        return 0
    except Exception as exc:
        print(f"ERROR: verification failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
