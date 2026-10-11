#!/usr/bin/env python3
"""Generate a workbook whose live formulas deliberately have no cached values."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SEED = 20250603
ROOT = Path(__file__).resolve().parent
FILES = ROOT / "files"
WORKBOOK_NAME = "seva_invoice_register_uncalculated.xlsx"


def make_clean_records() -> list[dict]:
    rng = np.random.default_rng(SEED)
    locations = [
        ("Community Nutrition", "Koraput", "Odisha"),
        ("Community Nutrition", "Rayagada", "Odisha"),
        ("School WASH", "Barmer", "Rajasthan"),
        ("School WASH", "Jaisalmer", "Rajasthan"),
        ("Women's Livelihoods", "Nandurbar", "Maharashtra"),
        ("Mobile Health Outreach", "Mandla", "Madhya Pradesh"),
        ("Mobile Health Outreach", "Dindori", "Madhya Pradesh"),
    ]
    funders = [
        ("Tata Trusts CSR", "CSR"),
        ("PM POSHAN", "Government Scheme"),
        ("Global Nutrition Fund (FCRA)", "Foreign Grant"),
        ("HDFC Bank Parivartan", "CSR"),
    ]
    narratives = [
        "Field nutrition kits",
        "Community facilitator supplies",
        "School water testing material",
        "Training venue and learning aids",
        "Mobile clinic consumables",
        "Livelihood demonstration inputs",
    ]
    records: list[dict] = []
    invoice_number = 1
    for month in (4, 5, 6):
        for within_month in range(12):
            programme, district, state = locations[(invoice_number * 3 + month) % len(locations)]
            funding_source, funding_type = funders[(invoice_number + month) % len(funders)]
            quantity = int(rng.integers(3, 18))
            unit_rate = int(rng.integers(12, 82)) * 1_000
            # Ensure the district ranking has a unique leader without hard-coding the answer.
            if district == "Koraput":
                unit_rate += 75_000
            records.append(
                {
                    "invoice_id": f"SEVA-25-{invoice_number:04d}",
                    "date": datetime(2025, month, 2 + ((within_month * 2 + month) % 25)),
                    "quantity": quantity,
                    "unit_rate": unit_rate,
                    "programme": programme,
                    "district": district,
                    "state": state,
                    "funding_source": funding_source,
                    "funding_type": funding_type,
                    "narrative": narratives[(invoice_number * 2 + month) % len(narratives)],
                }
            )
            invoice_number += 1
    return records


def truth_from_clean(records: list[dict]) -> dict:
    enriched = [{**row, "amount": row["quantity"] * row["unit_rate"]} for row in records]
    total = sum(row["amount"] for row in enriched)
    nutrition_odisha = sum(
        row["amount"]
        for row in enriched
        if row["programme"] == "Community Nutrition" and row["state"] == "Odisha"
    )
    government_count = len(
        {
            row["invoice_id"]
            for row in enriched
            if row["funding_type"] == "Government Scheme"
        }
    )
    by_district: dict[str, int] = {}
    for row in enriched:
        by_district[row["district"]] = by_district.get(row["district"], 0) + row["amount"]
    ranked = sorted(by_district.items(), key=lambda item: (-item[1], item[0]))
    if len(ranked) > 1 and ranked[0][1] == ranked[1][1]:
        raise RuntimeError("District ranking unexpectedly tied")
    foreign = sum(row["amount"] for row in enriched if row["funding_type"] == "Foreign Grant")
    return {
        "q1": total,
        "q2": nutrition_odisha,
        "q3": government_count,
        "q4": ranked[0][0],
        "q5": round(foreign / total * 100, 2),
    }


def write_case_json(truth: dict) -> None:
    case = {
        "id": "formulas_no_cache",
        "description": (
            "A quarterly NGO invoice register has live line-total, monthly-subtotal, "
            "and grand-total formulas but no cached formula results."
        ),
        "files": [WORKBOOK_NAME],
        "questions": [
            {
                "id": "q1",
                "question": (
                    "What is the total invoice spend in INR? Calculate each invoice line "
                    "as Quantity × Unit Rate, and ignore monthly subtotal and grand total rows."
                ),
                "answer": truth["q1"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "The live formula cells have no cached values and appear blank in data-only reads.",
            },
            {
                "id": "q2",
                "question": (
                    "What is the total computed spend in INR for Community Nutrition "
                    "in Odisha, using Quantity × Unit Rate for each invoice line?"
                ),
                "answer": truth["q2"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "Summing the uncached Line Amount column yields blanks instead of spend.",
            },
            {
                "id": "q3",
                "question": (
                    "How many distinct invoice IDs are funded by rows whose Funding Type "
                    "is exactly 'Government Scheme'? Count invoice detail rows only."
                ),
                "answer": truth["q3"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "Month labels and subtotal rows are interspersed with invoice rows.",
            },
            {
                "id": "q4",
                "question": (
                    "Which district has the highest total computed spend after calculating "
                    "Quantity × Unit Rate for every invoice?"
                ),
                "answer": truth["q4"],
                "kind": "name",
                "tol_pct": 0.0,
                "trap": "Ranking blank cached formula cells or the largest single invoice is wrong.",
            },
            {
                "id": "q5",
                "question": (
                    "What percentage of total computed invoice spend is funded by rows whose "
                    "Funding Type is 'Foreign Grant'? Use Quantity × Unit Rate and round to two decimals."
                ),
                "answer": truth["q5"],
                "kind": "number",
                "tol_pct": 0.01,
                "trap": "The percentage must be derived from inputs, not uncached formula results.",
            },
        ],
    }
    (ROOT / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def style_header(ws, row: int) -> None:
    thin = Side(style="thin", color="A6A6A6")
    for cell in ws[row][:11]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)


def write_workbook(records: list[dict]) -> None:
    FILES.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Register"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:K1")
    ws["A1"] = "SEVA COLLECTIVE — Q1 FY 2025-26 INVOICE REGISTER"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="375623")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:K2")
    ws["A2"] = (
        "Amounts in INR | Formula cells intentionally not calculated in this file; "
        "derive Line Amount as Quantity × Unit Rate"
    )
    ws["A2"].font = Font(italic=True, color="666666")
    headers = [
        "Invoice ID",
        "Expense Date",
        "Quantity",
        "Unit Rate (INR)",
        "Line Amount (formula)",
        "Programme",
        "District",
        "State",
        "Funding Source",
        "Funding Type",
        "Narrative",
    ]
    for col, value in enumerate(headers, start=1):
        ws.cell(4, col, value)
    style_header(ws, 4)

    month_names = {4: "APRIL 2025", 5: "MAY 2025", 6: "JUNE 2025"}
    subtotal_rows: list[int] = []
    for month in (4, 5, 6):
        month_heading_row = ws.max_row + 2
        ws.merge_cells(
            start_row=month_heading_row,
            start_column=1,
            end_row=month_heading_row,
            end_column=11,
        )
        ws.cell(month_heading_row, 1, month_names[month])
        ws.cell(month_heading_row, 1).font = Font(bold=True, color="1F4E78")
        ws.cell(month_heading_row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        first_detail_row = ws.max_row + 1
        for record in (row for row in records if row["date"].month == month):
            row_number = ws.max_row + 1
            ws.append(
                [
                    record["invoice_id"],
                    record["date"],
                    record["quantity"],
                    record["unit_rate"],
                    f"=C{row_number}*D{row_number}",
                    record["programme"],
                    record["district"],
                    record["state"],
                    record["funding_source"],
                    record["funding_type"],
                    record["narrative"],
                ]
            )
            ws.cell(row_number, 2).number_format = "dd-mmm-yyyy"
            ws.cell(row_number, 4).number_format = '₹#,##0;[Red]-₹#,##0'
            ws.cell(row_number, 5).number_format = '₹#,##0;[Red]-₹#,##0'
        last_detail_row = ws.max_row
        subtotal_row = last_detail_row + 1
        ws.cell(subtotal_row, 1, f"SUBTOTAL — {month_names[month]}")
        ws.cell(subtotal_row, 5, f"=SUM(E{first_detail_row}:E{last_detail_row})")
        for cell in ws[subtotal_row][:11]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
        ws.cell(subtotal_row, 5).number_format = '₹#,##0;[Red]-₹#,##0'
        subtotal_rows.append(subtotal_row)

    grand_total_row = ws.max_row + 2
    ws.cell(grand_total_row, 1, "GRAND TOTAL — Q1")
    subtotal_refs = ",".join(f"E{row}" for row in subtotal_rows)
    ws.cell(grand_total_row, 5, f"=SUM({subtotal_refs})")
    for cell in ws[grand_total_row][:11]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="375623")
    ws.cell(grand_total_row, 5).number_format = '₹#,##0;[Red]-₹#,##0'

    widths = [20, 15, 11, 19, 23, 28, 18, 20, 32, 22, 34]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{grand_total_row}"

    notes = wb.create_sheet("Formula notes")
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 105
    notes.append(["Item", "Instruction"])
    notes.append(["Line Amount", "Live formula = Quantity × Unit Rate. The workbook was not recalculated."])
    notes.append(["Monthly subtotal", "Live SUM formula over that month's line-amount cells."])
    notes.append(["Grand total", "Live SUM formula over the three monthly subtotal cells."])
    notes.append(["Analysis rule", "Compute from invoice inputs; exclude month labels, subtotals, and grand total."])
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")

    path = FILES / WORKBOOK_NAME
    wb.save(path)
    if path.stat().st_size >= 2 * 1024 * 1024:
        raise RuntimeError(f"{path.name} exceeds 2 MB")

    # openpyxl does not calculate formulas: the file must retain formulas while
    # exposing no cached results in data_only mode.
    formula_ws = load_workbook(path, data_only=False, read_only=True)["Invoice Register"]
    cached_ws = load_workbook(path, data_only=True, read_only=True)["Invoice Register"]
    invoice_rows = [
        row
        for row in range(1, formula_ws.max_row + 1)
        if str(formula_ws.cell(row, 1).value or "").startswith("SEVA-25-")
    ]
    if not invoice_rows:
        raise RuntimeError("No invoice rows were written")
    if not all(
        isinstance(formula_ws.cell(row, 5).value, str)
        and formula_ws.cell(row, 5).value.startswith("=")
        for row in invoice_rows
    ):
        raise RuntimeError("A line-total formula is missing")
    if not all(cached_ws.cell(row, 5).value is None for row in invoice_rows):
        raise RuntimeError("A line-total formula unexpectedly has a cached value")


def main() -> None:
    records = make_clean_records()
    truth = truth_from_clean(records)
    # The clean-record truth is finalized before any messy workbook rows are emitted.
    write_case_json(truth)
    write_workbook(records)
    print(f"generated formulas_no_cache: {len(records)} invoices, 5 questions")


if __name__ == "__main__":
    main()
