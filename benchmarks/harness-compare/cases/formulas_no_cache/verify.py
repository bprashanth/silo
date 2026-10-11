#!/usr/bin/env python3
"""Verify formulas_no_cache by calculating exclusively from invoice inputs."""

from __future__ import annotations

import json
import math
import posixpath
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
CASE_PATH = ROOT / "case.json"


def validate_uncached_formulas(path: Path) -> None:
    formula_ws = load_workbook(path, data_only=False, read_only=True)["Invoice Register"]
    cached_ws = load_workbook(path, data_only=True, read_only=True)["Invoice Register"]
    checked = 0
    line_formulas = 0
    subtotal_formulas = 0
    grand_formulas = 0
    for row_number in range(1, formula_ws.max_row + 1):
        label = str(formula_ws.cell(row_number, 1).value or "")
        if label.startswith("SEVA-25-"):
            category = "line"
        elif label.startswith("SUBTOTAL"):
            category = "subtotal"
        elif label.startswith("GRAND TOTAL"):
            category = "grand"
        else:
            continue
        formula = formula_ws.cell(row_number, 5).value
        cached = cached_ws.cell(row_number, 5).value
        if not isinstance(formula, str) or not formula.startswith("="):
            raise AssertionError(f"Row {row_number} is missing its live formula")
        if cached is not None:
            raise AssertionError(f"Row {row_number} unexpectedly has a cached formula value")
        if category == "line":
            if formula.replace("$", "") != f"=C{row_number}*D{row_number}":
                raise AssertionError(f"Row {row_number} has an unexpected line formula: {formula}")
            line_formulas += 1
        elif category == "subtotal":
            if not formula.upper().startswith("=SUM("):
                raise AssertionError(f"Row {row_number} subtotal is not a SUM formula")
            subtotal_formulas += 1
        else:
            if not formula.upper().startswith("=SUM("):
                raise AssertionError(f"Row {row_number} grand total is not a SUM formula")
            grand_formulas += 1
        checked += 1
    if line_formulas == 0 or subtotal_formulas != 3 or grand_formulas != 1:
        raise AssertionError(
            "Expected invoice formulas, three subtotal formulas, and one grand-total formula"
        )
    if checked != line_formulas + 4:
        raise AssertionError("Unexpected formula-row accounting")


def parse_with_openpyxl(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=False, read_only=True)["Invoice Register"]
    records: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        invoice_id = str(row[0] or "")
        if not invoice_id.startswith("SEVA-25-"):
            continue
        quantity = int(row[2])
        unit_rate = int(row[3])
        records.append(
            {
                "invoice_id": invoice_id,
                "quantity": quantity,
                "unit_rate": unit_rate,
                "programme": str(row[5]),
                "district": str(row[6]),
                "state": str(row[7]),
                "funding_type": str(row[9]),
            }
        )
    if not records:
        raise ValueError("Openpyxl parser found no invoice detail rows")
    return records


def derive_with_loops(records: list[dict]) -> dict:
    total = 0
    nutrition_odisha = 0
    government_ids: set[str] = set()
    district_totals: dict[str, int] = {}
    foreign = 0
    for row in records:
        amount = row["quantity"] * row["unit_rate"]
        total += amount
        if row["programme"] == "Community Nutrition" and row["state"] == "Odisha":
            nutrition_odisha += amount
        if row["funding_type"] == "Government Scheme":
            government_ids.add(row["invoice_id"])
        if row["funding_type"] == "Foreign Grant":
            foreign += amount
        district_totals[row["district"]] = district_totals.get(row["district"], 0) + amount
    ranking = sorted(district_totals.items(), key=lambda item: (-item[1], item[0]))
    if len(ranking) > 1 and ranking[0][1] == ranking[1][1]:
        raise ValueError("District ranking is tied")
    return {
        "q1": total,
        "q2": nutrition_odisha,
        "q3": len(government_ids),
        "q4": ranking[0][0],
        "q5": round(foreign / total * 100, 2),
    }


def read_sheet_xml(path: Path, sheet_name: str) -> dict[int, dict[int, object]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    doc_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet = next(
            node
            for node in workbook.findall(f".//{{{main_ns}}}sheet")
            if node.attrib["name"] == sheet_name
        )
        relation_id = sheet.attrib[f"{{{doc_rel_ns}}}id"]
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(
            node.attrib["Target"]
            for node in relationships.findall(f"{{{package_rel_ns}}}Relationship")
            if node.attrib["Id"] == relation_id
        )
        if target.startswith("/"):
            sheet_path = target.lstrip("/")
        elif target.startswith("xl/"):
            sheet_path = target
        else:
            sheet_path = posixpath.normpath(posixpath.join("xl", target))
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [
                "".join(text.text or "" for text in item.iter(f"{{{main_ns}}}t"))
                for item in shared_root.findall(f"{{{main_ns}}}si")
            ]
        root = ET.fromstring(archive.read(sheet_path))

    rows: dict[int, dict[int, object]] = {}
    for row_node in root.findall(f".//{{{main_ns}}}sheetData/{{{main_ns}}}row"):
        row_number = int(row_node.attrib["r"])
        values: dict[int, object] = {}
        for cell in row_node.findall(f"{{{main_ns}}}c"):
            reference = cell.attrib["r"]
            letters = "".join(char for char in reference if char.isalpha())
            column = 0
            for char in letters:
                column = column * 26 + ord(char.upper()) - ord("A") + 1
            cell_type = cell.attrib.get("t")
            value_node = cell.find(f"{{{main_ns}}}v")
            if cell_type == "inlineStr":
                inline = cell.find(f"{{{main_ns}}}is")
                value: object = (
                    "".join(text.text or "" for text in inline.iter(f"{{{main_ns}}}t"))
                    if inline is not None
                    else ""
                )
            elif cell_type == "s" and value_node is not None:
                value = shared[int(value_node.text)]
            elif value_node is None or value_node.text is None:
                value = None
            else:
                number = float(value_node.text)
                value = int(number) if number.is_integer() else number
            values[column] = value
        rows[row_number] = values
    return rows


def parse_and_derive_with_xml(path: Path) -> dict:
    rows = read_sheet_xml(path, "Invoice Register")
    details = [
        row
        for row in rows.values()
        if str(row.get(1, "")).startswith("SEVA-25-")
    ]
    if not details:
        raise ValueError("OOXML parser found no invoice detail rows")
    computed = [(row, int(row[3]) * int(row[4])) for row in details]
    district_totals: dict[str, int] = defaultdict(int)
    for row, amount in computed:
        district_totals[str(row[7])] += amount
    ranking = sorted(district_totals.items(), key=lambda item: (-item[1], item[0]))
    if len(ranking) > 1 and ranking[0][1] == ranking[1][1]:
        raise ValueError("District ranking is tied")
    total = sum(amount for _, amount in computed)
    foreign = sum(
        amount for row, amount in computed if row[10] == "Foreign Grant"
    )
    return {
        "q1": total,
        "q2": sum(
            amount
            for row, amount in computed
            if row[6] == "Community Nutrition" and row[8] == "Odisha"
        ),
        "q3": len(
            {
                str(row[1])
                for row, _ in computed
                if row[10] == "Government Scheme"
            }
        ),
        "q4": ranking[0][0],
        "q5": round(foreign / total * 100, 2),
    }


def values_match(left, right, tol_pct: float = 0.0) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        return left == right
    allowed = max(1e-9, abs(float(right)) * tol_pct / 100)
    return math.isclose(float(left), float(right), rel_tol=0.0, abs_tol=allowed)


def main() -> int:
    try:
        case = json.loads(CASE_PATH.read_text(encoding="utf-8"))
        if case.get("id") != "formulas_no_cache":
            raise ValueError("case.json has the wrong case id")
        if len(case.get("questions", [])) not in (5, 6):
            raise ValueError("case.json must contain 5 or 6 questions")
        if len(case.get("files", [])) != 1:
            raise ValueError("case.json must name exactly one workbook")
        workbook = ROOT / "files" / case["files"][0]
        if not workbook.is_file() or workbook.stat().st_size >= 2 * 1024 * 1024:
            raise ValueError("Workbook is missing or is not under 2 MB")

        validate_uncached_formulas(workbook)
        first = derive_with_loops(parse_with_openpyxl(workbook))
        second = parse_and_derive_with_xml(workbook)
        if set(first) != set(second):
            raise AssertionError("Independent derivations produced different question sets")
        for question_id in first:
            if not values_match(first[question_id], second[question_id]):
                raise AssertionError(
                    f"{question_id}: independent derivations disagree: "
                    f"{first[question_id]!r} != {second[question_id]!r}"
                )

        questions = case["questions"]
        expected_ids = {question["id"] for question in questions}
        if expected_ids != set(first) or len(expected_ids) != len(questions):
            raise AssertionError("Question IDs are missing, duplicated, or unexpected")
        for question in questions:
            actual = first[question["id"]]
            if not values_match(actual, question["answer"], float(question["tol_pct"])):
                raise AssertionError(
                    f"{question['id']}: expected {question['answer']!r}, got {actual!r}"
                )
        print(
            f"PASS formulas_no_cache: {len(questions)} questions; "
            "all formula caches empty; openpyxl and raw OOXML derivations agree"
        )
        return 0
    except Exception as exc:
        print(f"FAIL formulas_no_cache: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
