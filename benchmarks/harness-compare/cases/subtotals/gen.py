#!/usr/bin/env python3
"""Generate the deterministic subtotals benchmark case."""

from __future__ import annotations

import json
import random
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SEED = 20250417
HERE = Path(__file__).resolve().parent
FILES = HERE / "files"
OUTPUT = FILES / "monthly_expense_ledger.xlsx"

MONTHS = [
    (4, "Apr"),
    (5, "May"),
    (6, "Jun"),
    (7, "Jul"),
    (8, "Aug"),
    (9, "Sep"),
]

PROGRAMMES = [
    ("Jal Jeevan Saathi", "Nashik", "Maharashtra", 1.55),
    ("Poshan Shakti", "Nandurbar", "Maharashtra", 1.12),
    ("Shiksha Udaan", "Barmer", "Rajasthan", 0.93),
    ("Swasthya Sathi", "Koraput", "Odisha", 1.28),
    ("Mahila Ajeevika", "Gaya", "Bihar", 0.88),
    ("Krishi Pragati", "Kalahandi", "Odisha", 1.04),
]

DONORS = [
    ("Azim Premji Foundation", 0.22),
    ("Tata Trusts CSR", 0.19),
    ("Jal Jeevan Mission", 0.20),
    ("HDFC Bank Parivartan", 0.16),
    ("Global Green Fund", 0.13),
    ("State Government Grant", 0.10),
]

EXPENSE_HEADS = [
    ("Field staff and facilitators", 1.35),
    ("Training and community meetings", 0.82),
    ("Programme materials", 1.08),
    ("Travel and local conveyance", 0.64),
    ("Monitoring and surveys", 0.76),
    ("Equipment and minor works", 1.52),
]


def make_records() -> list[dict]:
    """Build clean records; all truth is derived from this list."""
    random.seed(SEED)
    rng = np.random.default_rng(SEED)
    records: list[dict] = []
    serial = 1

    for month_num, _month_label in MONTHS:
        month_start = date(2025, month_num, 1)
        next_month = date(2025, month_num + 1, 1) if month_num < 12 else date(2026, 1, 1)
        days = (next_month - month_start).days
        # Different block sizes make the monthly subtotals look organic.
        rows_this_month = 15 + (month_num % 3)
        for row_index in range(rows_this_month):
            programme = PROGRAMMES[(row_index + month_num + int(rng.integers(0, 4))) % len(PROGRAMMES)]
            donor_index = int(rng.choice(len(DONORS), p=[d[1] for d in DONORS]))
            donor = DONORS[donor_index][0]
            head_name, head_weight = EXPENSE_HEADS[int(rng.integers(0, len(EXPENSE_HEADS)))]
            amount = int(
                round(
                    rng.integers(42_000, 238_000)
                    * programme[3]
                    * head_weight
                    / 100
                )
                * 100
            )
            records.append(
                {
                    "transaction_id": f"EXP-25-{serial:04d}",
                    "date": month_start + timedelta(days=int(rng.integers(0, days))),
                    "programme": programme[0],
                    "district": programme[1],
                    "state": programme[2],
                    "donor": donor,
                    "expense_head": head_name,
                    "amount": amount,
                    "remarks": random.choice(
                        [
                            "Approved against monthly plan",
                            "Supporting vouchers filed",
                            "District team reimbursement",
                            "Direct vendor payment",
                            "",
                        ]
                    ),
                }
            )
            serial += 1

    return records


def answers_from_clean(records: list[dict]) -> dict[str, int | float | str]:
    total = sum(r["amount"] for r in records)
    water_maharashtra = sum(
        r["amount"]
        for r in records
        if r["programme"] == "Jal Jeevan Saathi" and r["state"] == "Maharashtra"
    )
    poshan_count = sum(r["programme"] == "Poshan Shakti" for r in records)

    district_totals: dict[str, int] = {}
    programme_totals: dict[str, int] = {}
    for record in records:
        district_totals[record["district"]] = district_totals.get(record["district"], 0) + record["amount"]
        programme_totals[record["programme"]] = programme_totals.get(record["programme"], 0) + record["amount"]

    top_district = max(district_totals.items(), key=lambda item: (item[1], item[0]))[0]
    top_programme = max(programme_totals.items(), key=lambda item: (item[1], item[0]))[0]
    azim_total = sum(r["amount"] for r in records if r["donor"] == "Azim Premji Foundation")

    return {
        "q1": total,
        "q2": water_maharashtra,
        "q3": poshan_count,
        "q4": top_district,
        "q5": round(100.0 * azim_total / total, 2),
        "q6": top_programme,
    }


def write_workbook(records: list[dict]) -> None:
    if FILES.exists():
        shutil.rmtree(FILES)
    FILES.mkdir(parents=True)

    wb = Workbook()
    wb.properties.creator = "SevaSetu Foundation"
    wb.properties.created = datetime(2025, 10, 3, 10, 0, 0)
    wb.properties.modified = datetime(2025, 10, 3, 10, 0, 0)
    ws = wb.active
    ws.title = "Expense Ledger"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    ws["A1"] = "SevaSetu Foundation — Monthly Expense Ledger"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:I2")
    ws["A2"] = "FY 2025–26 | Actual expenditure | Amounts in INR"
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "Transaction ID",
        "Date",
        "Programme",
        "District",
        "State",
        "Funding Source",
        "Expense Head",
        "Amount (INR)",
        "Remarks",
    ]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for month_num, month_label in MONTHS:
        ws.append([f"{date(2025, month_num, 1):%B %Y}".upper()])
        month_heading_row = ws.max_row
        ws.merge_cells(start_row=month_heading_row, start_column=1, end_row=month_heading_row, end_column=9)
        ws.cell(month_heading_row, 1).font = Font(bold=True, color="1F1F1F")
        ws.cell(month_heading_row, 1).fill = PatternFill("solid", fgColor="D9EAF7")

        month_records = [r for r in records if r["date"].month == month_num]
        for record in month_records:
            ws.append(
                [
                    record["transaction_id"],
                    record["date"],
                    record["programme"],
                    record["district"],
                    record["state"],
                    record["donor"],
                    record["expense_head"],
                    record["amount"],
                    record["remarks"],
                ]
            )
            ws.cell(ws.max_row, 2).number_format = "dd-mmm-yyyy"
            ws.cell(ws.max_row, 8).number_format = '₹#,##0;[Red]-₹#,##0'

        ws.append([f"Total – {month_label} 2025", None, None, None, None, None, None, sum(r["amount"] for r in month_records)])
        subtotal_row = ws.max_row
        ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=7)
        for cell in ws[subtotal_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(subtotal_row, 1).alignment = Alignment(horizontal="right")
        ws.cell(subtotal_row, 8).number_format = '₹#,##0;[Red]-₹#,##0'
        ws.append([])

    grand_total = sum(r["amount"] for r in records)
    ws.append(["GRAND TOTAL — APR TO SEP 2025", None, None, None, None, None, None, grand_total])
    grand_row = ws.max_row
    ws.merge_cells(start_row=grand_row, start_column=1, end_row=grand_row, end_column=7)
    for cell in ws[grand_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.cell(grand_row, 1).alignment = Alignment(horizontal="right")
    ws.cell(grand_row, 8).number_format = '₹#,##0;[Red]-₹#,##0'

    widths = [18, 14, 23, 15, 16, 27, 31, 17, 29]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{grand_row}"
    wb.save(OUTPUT)


def write_case_json(answers: dict[str, int | float | str]) -> None:
    questions = [
        {
            "id": "q1",
            "question": "Excluding every monthly subtotal and the grand total row, what was total actual expenditure from April through September 2025, in INR?",
            "answer": answers["q1"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Adding subtotal or grand-total rows double or triple counts the same expenses.",
        },
        {
            "id": "q2",
            "question": "Excluding subtotal and grand-total rows, how much was spent on Jal Jeevan Saathi in Maharashtra from April through September 2025, in INR?",
            "answer": answers["q2"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "The ledger's summary rows have blank programme and state fields and must not enter the filter.",
        },
        {
            "id": "q3",
            "question": "How many underlying expense transactions (rows with an EXP- transaction ID, not subtotal rows) were recorded for Poshan Shakti?",
            "answer": answers["q3"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Counting visual month blocks or summary rows instead of transaction rows gives the wrong count.",
        },
        {
            "id": "q4",
            "question": "After excluding all subtotal and grand-total rows, which district had the highest total actual expenditure?",
            "answer": answers["q4"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "Summing the precomputed totals alongside detail rows corrupts district rankings.",
        },
        {
            "id": "q5",
            "question": "What percentage of total underlying expenditure was funded by Azim Premji Foundation? Exclude all subtotal and grand-total rows and report the percentage to two decimal places.",
            "answer": answers["q5"],
            "kind": "number",
            "tol_pct": 0.01,
            "trap": "Using a denominator that includes monthly or grand totals understates the donor share.",
        },
        {
            "id": "q6",
            "question": "After excluding monthly subtotal and grand-total rows, which programme had the highest total actual expenditure?",
            "answer": answers["q6"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "The bold totals are not programme records and must not be grouped as if they were.",
        },
    ]
    payload = {
        "id": "subtotals",
        "description": "A six-month NGO expense ledger with monthly subtotal rows, blank separators, and a grand total designed to tempt double counting.",
        "files": [OUTPUT.name],
        "questions": questions,
    }
    (HERE / "case.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main() -> None:
    records = make_records()
    answers = answers_from_clean(records)
    write_workbook(records)
    write_case_json(answers)
    print(f"subtotals: generated {len(records)} clean transactions and {len(answers)} questions")


if __name__ == "__main__":
    main()
