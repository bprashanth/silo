#!/usr/bin/env python3
"""Independently verify every answer in the subtotals case."""

from __future__ import annotations

import json
import math
import re
import sys
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
CASE_PATH = HERE / "case.json"
MONTH_HEADING = re.compile(
    r"^(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER) 2025$"
)
TXN_ID = re.compile(r"^EXP-\d{2}-\d{4}$")


def fail(message: str) -> None:
    raise AssertionError(message)


def amount(value: object) -> int:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        fail("missing amount in transaction row")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        fail(f"non-numeric amount in transaction row: {value!r}")
    if not float(value).is_integer():
        fail(f"transaction amount is not whole INR: {value!r}")
    return int(value)


def parsed_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%d-%b-%Y"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                pass
    fail(f"unparseable transaction date: {value!r}")


def read_sheet_xml(workbook: Path) -> list[list[object]]:
    """Read sheet1 directly from XLSX XML without using openpyxl's cell model."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(workbook) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))

    cells: dict[tuple[int, int], object] = {}
    max_row = max_col = 0
    for cell in root.iter(f"{{{main_ns}}}c"):
        reference = cell.attrib["r"]
        match = re.fullmatch(r"([A-Z]+)(\d+)", reference)
        if not match:
            fail(f"invalid cell reference in worksheet XML: {reference!r}")
        letters, row_text = match.groups()
        column = 0
        for character in letters:
            column = column * 26 + ord(character) - ord("A") + 1
        row = int(row_text)
        cell_type = cell.attrib.get("t", "n")
        if cell_type == "inlineStr":
            inline = cell.find(f"{{{main_ns}}}is")
            value: object = "" if inline is None else "".join(
                node.text or "" for node in inline.iter(f"{{{main_ns}}}t")
            )
        else:
            value_node = cell.find(f"{{{main_ns}}}v")
            raw = None if value_node is None else value_node.text
            if raw is None:
                value = None
            elif cell_type == "s":
                value = shared[int(raw)]
            elif cell_type == "b":
                value = raw == "1"
            elif cell_type in ("str", "e"):
                value = raw
            else:
                number = float(raw)
                value = int(number) if number.is_integer() else number
        cells[(row, column)] = value
        max_row = max(max_row, row)
        max_col = max(max_col, column)
    return [
        [cells.get((row, column)) for column in range(1, max_col + 1)]
        for row in range(1, max_row + 1)
    ]


def normalized_record(values: list[object], columns: dict[str, int]) -> dict:
    transaction_id = str(values[columns["Transaction ID"]]).strip()
    if not TXN_ID.fullmatch(transaction_id):
        fail(f"bad transaction ID: {transaction_id!r}")
    raw_date = values[columns["Date"]]
    record = {
        "transaction_id": transaction_id,
        "date": parsed_date(raw_date),
        "programme": str(values[columns["Programme"]]).strip(),
        "district": str(values[columns["District"]]).strip(),
        "state": str(values[columns["State"]]).strip(),
        "donor": str(values[columns["Funding Source"]]).strip(),
        "expense_head": str(values[columns["Expense Head"]]).strip(),
        "amount": amount(values[columns["Amount (INR)"]]),
    }
    for field in ("programme", "district", "state", "donor", "expense_head"):
        if not record[field] or record[field] == "None":
            fail(f"blank {field} for {transaction_id}")
    return record


def parse_by_transaction_id(workbook: Path) -> list[dict]:
    """Derivation A: discover columns, then select only explicit transaction IDs."""
    wb = load_workbook(workbook, data_only=False, read_only=True)
    if wb.sheetnames != ["Expense Ledger"]:
        fail(f"unexpected sheets: {wb.sheetnames}")
    ws = wb["Expense Ledger"]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    required = {
        "Transaction ID",
        "Date",
        "Programme",
        "District",
        "State",
        "Funding Source",
        "Expense Head",
        "Amount (INR)",
    }
    header_indexes = [i for i, row in enumerate(rows) if required.issubset(set(row))]
    if len(header_indexes) != 1:
        fail(f"expected one ledger header, found {len(header_indexes)}")
    header_index = header_indexes[0]
    columns = {str(value): index for index, value in enumerate(rows[header_index]) if value is not None}

    records = []
    for row in rows[header_index + 1 :]:
        first = row[columns["Transaction ID"]]
        if isinstance(first, str) and TXN_ID.fullmatch(first.strip()):
            records.append(normalized_record(row, columns))
    if not records:
        fail("no transaction records found")
    if len({r["transaction_id"] for r in records}) != len(records):
        fail("duplicate transaction IDs")
    return records


def parse_by_month_blocks(workbook: Path) -> list[dict]:
    """Derivation B: parse raw XLSX XML and use structural month boundaries."""
    rows = read_sheet_xml(workbook)
    required = [
        "Transaction ID",
        "Date",
        "Programme",
        "District",
        "State",
        "Funding Source",
        "Expense Head",
        "Amount (INR)",
    ]
    header_indexes = [
        index for index, row in enumerate(rows) if all(label in row for label in required)
    ]
    if len(header_indexes) != 1:
        fail(f"block parser found {len(header_indexes)} headers")
    columns = {label: rows[header_indexes[0]].index(label) for label in required}

    block_starts = [
        index
        for index, row in enumerate(rows)
        if isinstance(row[0], str) and MONTH_HEADING.fullmatch(row[0].strip())
    ]
    if len(block_starts) != 6:
        fail(f"expected 6 month blocks, found {len(block_starts)}")

    records: list[dict] = []
    claimed_subtotals: list[int] = []
    calculated_subtotals: list[int] = []
    for start in block_starts:
        block_records: list[dict] = []
        subtotal_seen = False
        for row in rows[start + 1 :]:
            label = row[0]
            if isinstance(label, str) and label.startswith("Total"):
                claimed_subtotals.append(amount(row[columns["Amount (INR)"]]))
                calculated_subtotals.append(sum(record["amount"] for record in block_records))
                subtotal_seen = True
                break
            # A real detail row has a parseable date; blank separator rows do not.
            raw_date = row[columns["Date"]]
            if raw_date is None:
                fail(f"unexpected non-detail row inside block beginning at spreadsheet row {start + 1}")
            parsed_date(raw_date)
            block_records.append(normalized_record(row, columns))
        if not subtotal_seen:
            fail(f"month block beginning at spreadsheet row {start + 1} lacks subtotal")
        records.extend(block_records)

    if claimed_subtotals != calculated_subtotals:
        fail(f"monthly subtotal mismatch: claimed={claimed_subtotals}, calculated={calculated_subtotals}")

    grand_rows = [
        row for row in rows if isinstance(row[0], str) and row[0].strip().startswith("GRAND TOTAL")
    ]
    if len(grand_rows) != 1:
        fail(f"expected one grand total, found {len(grand_rows)}")
    claimed_grand = amount(grand_rows[0][columns["Amount (INR)"]])
    if claimed_grand != sum(calculated_subtotals):
        fail(f"grand total {claimed_grand} disagrees with month detail {sum(calculated_subtotals)}")
    return records


def derive_with_loops(records: list[dict]) -> dict[str, int | float | str]:
    total = sum(record["amount"] for record in records)
    district_totals: dict[str, int] = defaultdict(int)
    programme_totals: dict[str, int] = defaultdict(int)
    for record in records:
        district_totals[record["district"]] += record["amount"]
        programme_totals[record["programme"]] += record["amount"]
    return {
        "q1": total,
        "q2": sum(
            record["amount"]
            for record in records
            if record["programme"] == "Jal Jeevan Saathi" and record["state"] == "Maharashtra"
        ),
        "q3": sum(record["programme"] == "Poshan Shakti" for record in records),
        "q4": max(district_totals.items(), key=lambda item: (item[1], item[0]))[0],
        "q5": round(
            100.0
            * sum(record["amount"] for record in records if record["donor"] == "Azim Premji Foundation")
            / total,
            2,
        ),
        "q6": max(programme_totals.items(), key=lambda item: (item[1], item[0]))[0],
    }


def derive_with_indexes(records: list[dict]) -> dict[str, int | float | str]:
    amounts = [record["amount"] for record in records]
    total = sum(amounts)
    district_totals = {
        district: sum(record["amount"] for record in records if record["district"] == district)
        for district in {record["district"] for record in records}
    }
    programme_totals = {
        programme: sum(record["amount"] for record in records if record["programme"] == programme)
        for programme in {record["programme"] for record in records}
    }
    return {
        "q1": total,
        "q2": sum(
            record["amount"]
            for record in records
            if record["programme"] == "Jal Jeevan Saathi" and record["state"] == "Maharashtra"
        ),
        "q3": len([record for record in records if record["programme"] == "Poshan Shakti"]),
        "q4": max(district_totals.items(), key=lambda item: (item[1], item[0]))[0],
        "q5": round(
            100.0
            * sum(
                record["amount"]
                for record in records
                if record["donor"] == "Azim Premji Foundation"
            )
            / total,
            2,
        ),
        "q6": max(programme_totals.items(), key=lambda item: (item[1], item[0]))[0],
    }


def values_agree(left: object, right: object, kind: str, tol_pct: float) -> bool:
    if kind == "name":
        return isinstance(left, str) and isinstance(right, str) and left == right
    if isinstance(left, bool) or isinstance(right, bool):
        return False
    try:
        left_number = float(left)
        right_number = float(right)
    except (TypeError, ValueError):
        return False
    tolerance = abs(right_number) * tol_pct / 100.0
    return abs(left_number - right_number) <= max(tolerance, 1e-9)


def main() -> None:
    case = json.loads(CASE_PATH.read_text(encoding="utf-8"))
    if case.get("id") != "subtotals":
        fail("case id is not subtotals")
    questions = case.get("questions")
    if not isinstance(questions, list) or len(questions) not in (5, 6):
        fail("case must contain 5 or 6 questions")
    if len({question.get("id") for question in questions}) != len(questions):
        fail("question IDs are not unique")
    files = case.get("files")
    if files != ["monthly_expense_ledger.xlsx"]:
        fail(f"unexpected file manifest: {files!r}")
    workbook = HERE / "files" / files[0]
    if not workbook.is_file():
        fail(f"missing workbook: {workbook}")
    if workbook.stat().st_size >= 2_000_000:
        fail(f"workbook is not under 2 MB: {workbook.stat().st_size} bytes")

    records_a = parse_by_transaction_id(workbook)
    records_b = parse_by_month_blocks(workbook)
    canonical = lambda records: sorted(
        (
            record["transaction_id"],
            record["date"].isoformat(),
            record["programme"],
            record["district"],
            record["state"],
            record["donor"],
            record["expense_head"],
            record["amount"],
        )
        for record in records
    )
    if canonical(records_a) != canonical(records_b):
        fail("independent parsers did not recover the same transaction records")
    answers_a = derive_with_loops(records_a)
    answers_b = derive_with_indexes(records_b)

    if set(answers_a) != {question["id"] for question in questions}:
        fail("derived answer IDs do not exactly match case questions")
    for question in questions:
        question_id = question["id"]
        kind = question.get("kind")
        tolerance = float(question.get("tol_pct", 0.0))
        if kind not in ("number", "name"):
            fail(f"{question_id}: invalid kind {kind!r}")
        if not question.get("question") or not question.get("trap"):
            fail(f"{question_id}: missing question or trap")
        if not values_agree(answers_a[question_id], answers_b[question_id], kind, 0.0):
            fail(
                f"{question_id}: independent derivations disagree: "
                f"{answers_a[question_id]!r} != {answers_b[question_id]!r}"
            )
        if not values_agree(answers_a[question_id], question.get("answer"), kind, tolerance):
            fail(
                f"{question_id}: stored answer {question.get('answer')!r} "
                f"does not match derived {answers_a[question_id]!r}"
            )

    print(
        f"subtotals: PASS — {len(records_a)} transactions, "
        f"{len(questions)} questions, 2 independent derivations agree"
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"subtotals: FAIL — {exc}", file=sys.stderr)
        raise SystemExit(1)
