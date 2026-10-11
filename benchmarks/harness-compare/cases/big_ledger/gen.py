#!/usr/bin/env python3
"""Generate the deterministic big_ledger benchmark case."""

from __future__ import annotations

import json
import random
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SEED = 20250917
FIELD_ROWS = 12_000
SALARY_ROWS = 8_000
ROOT = Path(__file__).resolve().parent
FILES = ROOT / "files"
WORKBOOK_NAME = "national_programme_ledger.xlsx"

LOOKUP = [
    ("CC-MCH-OD-GAN", "Maternal and Child Health", "Ganjam", "Odisha", "National Health Mission"),
    ("CC-MCH-OD-KOR", "Maternal and Child Health", "Koraput", "Odisha", "National Health Mission"),
    ("CC-MCH-KA-RAI", "Maternal and Child Health", "Raichur", "Karnataka", "Tata CSR Trust"),
    ("CC-MCH-JH-DUM", "Maternal and Child Health", "Dumka", "Jharkhand", "Global Health Fund"),
    ("CC-RL-OD-MAY", "Rural Livelihoods", "Mayurbhanj", "Odisha", "Deendayal Antyodaya Yojana"),
    ("CC-RL-KA-KOP", "Rural Livelihoods", "Koppal", "Karnataka", "HDFC Parivartan"),
    ("CC-RL-RJ-BAR", "Rural Livelihoods", "Barmer", "Rajasthan", "Rajasthan Grameen Mission"),
    ("CC-RL-JH-DUM", "Rural Livelihoods", "Dumka", "Jharkhand", "Ford Foundation"),
    ("CC-TE-OD-KOR", "Tribal Education", "Koraput", "Odisha", "Samagra Shiksha"),
    ("CC-TE-OD-MAY", "Tribal Education", "Mayurbhanj", "Odisha", "Azim Premji Foundation"),
    ("CC-TE-RJ-BAR", "Tribal Education", "Barmer", "Rajasthan", "Educate India CSR Trust"),
    ("CC-TE-JH-WES", "Tribal Education", "West Singhbhum", "Jharkhand", "Samagra Shiksha"),
    ("CC-WS-OD-GAN", "Water and Sanitation", "Ganjam", "Odisha", "Jal Jeevan Mission"),
    ("CC-WS-KA-RAI", "Water and Sanitation", "Raichur", "Karnataka", "Infosys Foundation"),
    ("CC-WS-RJ-BAR", "Water and Sanitation", "Barmer", "Rajasthan", "Jal Jeevan Mission"),
    ("CC-WS-JH-WES", "Water and Sanitation", "West Singhbhum", "Jharkhand", "WaterAid India"),
    ("CC-NS-OD-GAN", "Nutrition Security", "Ganjam", "Odisha", "POSHAN Abhiyaan"),
    ("CC-NS-KA-KOP", "Nutrition Security", "Koppal", "Karnataka", "Tata CSR Trust"),
    ("CC-NS-RJ-BAR", "Nutrition Security", "Barmer", "Rajasthan", "World Food Programme"),
    ("CC-NS-JH-DUM", "Nutrition Security", "Dumka", "Jharkhand", "POSHAN Abhiyaan"),
    ("CC-CR-OD-KOR", "Climate Resilience", "Koraput", "Odisha", "Green Climate Fund"),
    ("CC-CR-KA-RAI", "Climate Resilience", "Raichur", "Karnataka", "SELCO Foundation"),
    ("CC-CR-RJ-BAR", "Climate Resilience", "Barmer", "Rajasthan", "Green Climate Fund"),
    ("CC-CR-JH-WES", "Climate Resilience", "West Singhbhum", "Jharkhand", "Oak Foundation"),
]

PROGRAMME_FACTOR = {
    "Maternal and Child Health": 1.10,
    "Rural Livelihoods": 1.25,
    "Tribal Education": 0.90,
    "Water and Sanitation": 1.45,
    "Nutrition Security": 1.00,
    "Climate Resilience": 1.18,
}


def make_clean_records() -> tuple[list[dict], list[dict]]:
    rng = random.Random(SEED)
    lookup_by_code = {
        code: {
            "programme": programme,
            "district": district,
            "state": state,
            "donor": donor,
        }
        for code, programme, district, state, donor in LOOKUP
    }
    codes = list(lookup_by_code)
    expense_heads = {
        "Field travel": (800, 28_000),
        "Community training": (8_000, 160_000),
        "Medical supplies": (4_000, 220_000),
        "School materials": (3_000, 140_000),
        "Water infrastructure": (20_000, 480_000),
        "Monitoring visit": (2_000, 55_000),
    }
    vendors = [
        "Maa Tarini Enterprises",
        "Kaveri Logistics",
        "Jai Marwar Supplies",
        "Adivasi Seva Cooperative",
        "Bharat Training Associates",
        "District Resource Centre",
    ]
    field_records = []
    start = date(2025, 4, 1)
    for index in range(FIELD_ROWS):
        unknown = rng.random() < 0.015
        code = f"CC-UNM-{rng.randrange(1, 10):03d}" if unknown else rng.choice(codes)
        expense_head = rng.choice(list(expense_heads))
        low, high = expense_heads[expense_head]
        programme = lookup_by_code.get(code, {}).get("programme")
        factor = PROGRAMME_FACTOR.get(programme, 1.0)
        amount = int(round(rng.randrange(low, high + 1) * factor / 100.0) * 100)
        status_roll = rng.random()
        status = "Posted" if status_roll < 0.91 else ("Void" if status_roll < 0.96 else "Draft")
        field_records.append(
            {
                "id": f"FE-{index + 1:06d}",
                "date": start + timedelta(days=rng.randrange(365)),
                "code": code,
                "expense_head": expense_head,
                "vendor": rng.choice(vendors),
                "amount": amount,
                "status": status,
            }
        )

    roles = {
        "Community Mobiliser": (24_000, 38_000),
        "Programme Officer": (42_000, 68_000),
        "District Coordinator": (58_000, 88_000),
        "Nurse Counsellor": (35_000, 62_000),
        "Monitoring Specialist": (55_000, 96_000),
        "Programme Manager": (82_000, 135_000),
    }
    salary_records = []
    for index in range(SALARY_ROWS):
        unknown = rng.random() < 0.012
        code = f"CC-UNM-{rng.randrange(1, 10):03d}" if unknown else rng.choice(codes)
        role = rng.choice(list(roles))
        low, high = roles[role]
        programme = lookup_by_code.get(code, {}).get("programme")
        factor = 0.92 + (PROGRAMME_FACTOR.get(programme, 1.0) - 1.0) * 0.35
        amount = int(round(rng.randrange(low, high + 1) * factor / 100.0) * 100)
        status_roll = rng.random()
        status = "Paid" if status_roll < 0.94 else ("Held" if status_roll < 0.975 else "Reversed")
        month_offset = rng.randrange(12)
        salary_records.append(
            {
                "id": f"PAY-{index + 1:06d}",
                "month": date(2025 + (3 + month_offset) // 12, (3 + month_offset) % 12 + 1, 1),
                "employee_id": f"EMP-{rng.randrange(1, 901):04d}",
                "role": role,
                "code": code,
                "amount": amount,
                "status": status,
            }
        )
    return field_records, salary_records


def compute_truth(field_records: list[dict], salary_records: list[dict]) -> dict:
    lookup = {
        code: {
            "programme": programme,
            "district": district,
            "state": state,
        }
        for code, programme, district, state, _ in LOOKUP
    }
    joined = []
    for source, rows, included_status in (
        ("field", field_records, "Posted"),
        ("salary", salary_records, "Paid"),
    ):
        for row in rows:
            mapped = lookup.get(row["code"])
            if mapped is None or row["status"] != included_status:
                continue
            joined.append(
                {
                    "source": source,
                    "amount": row["amount"],
                    **mapped,
                }
            )

    total = sum(row["amount"] for row in joined)
    q2 = sum(
        row["amount"]
        for row in joined
        if row["programme"] == "Maternal and Child Health"
        and row["state"] == "Odisha"
    )
    q3 = sum(row["district"] == "Barmer" for row in joined)
    programme_totals = defaultdict(int)
    district_totals = defaultdict(int)
    for row in joined:
        programme_totals[row["programme"]] += row["amount"]
        district_totals[row["district"]] += row["amount"]
    q4 = max(programme_totals, key=programme_totals.get)
    rural = [row for row in joined if row["programme"] == "Rural Livelihoods"]
    q5 = 100.0 * sum(
        row["amount"] for row in rural if row["source"] == "salary"
    ) / sum(row["amount"] for row in rural)
    q6 = max(district_totals, key=district_totals.get)

    assert list(programme_totals.values()).count(programme_totals[q4]) == 1
    assert list(district_totals.values()).count(district_totals[q6]) == 1
    return {
        "q1": total,
        "q2": q2,
        "q3": q3,
        "q4": q4,
        "q5": round(q5, 6),
        "q6": q6,
    }


def messy_code(code: str, index: int) -> str:
    if index % 17 == 0:
        return f" {code.lower()} "
    if index % 11 == 0:
        return f"{code} "
    if index % 7 == 0:
        return code.lower()
    return code


def messy_status(status: str, index: int) -> str:
    if index % 19 == 0:
        return f" {status.lower()} "
    if index % 13 == 0:
        return status.upper()
    return status


def indian_currency(amount: int) -> str:
    digits = str(amount)
    if len(digits) <= 3:
        grouped = digits
    else:
        last = digits[-3:]
        prefix = digits[:-3]
        groups = []
        while prefix:
            groups.append(prefix[-2:])
            prefix = prefix[:-2]
        grouped = ",".join(reversed(groups)) + "," + last
    return "₹" + grouped


def display_amount(amount: int, index: int) -> int | str:
    if index % 9 == 0:
        return indian_currency(amount)
    if index % 14 == 0:
        return f"{amount:,}"
    return amount


def style_header(ws, row_number: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row_number]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def write_workbook(field_records: list[dict], salary_records: list[dict]) -> None:
    FILES.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    field = wb.active
    field.title = "Field Expenses"
    field.merge_cells("A1:G1")
    field["A1"] = "Sahyog Foundation — National Field Expense Ledger FY 2025-26"
    field["A1"].font = Font(bold=True, size=14)
    field["A2"] = (
        "Use only transaction IDs beginning FE-. Posting status and cost-centre "
        "codes may have inconsistent case or surrounding spaces."
    )
    field.merge_cells("A2:G2")
    field["A3"] = "Prepared by Finance Operations"
    field.append([])
    field_headers = [
        "Transaction ID",
        "Transaction Date",
        "Cost Centre",
        "Expense Head",
        "Vendor / Payee",
        "Amount (INR)",
        "Posting Status",
    ]
    field.append(field_headers)
    style_header(field, 5)
    batch_total = 0
    for index, row in enumerate(field_records, start=1):
        field.append(
            [
                row["id"],
                row["date"],
                messy_code(row["code"], index),
                row["expense_head"],
                row["vendor"],
                display_amount(row["amount"], index),
                messy_status(row["status"], index),
            ]
        )
        batch_total += row["amount"]
        if index % 2_500 == 0 and index != len(field_records):
            field.append([f"Batch {index // 2500} subtotal — do not post", None, None, None, None, batch_total, None])
            field.append([])
            field.append(field_headers)
            batch_total = 0

    salary = wb.create_sheet("Salary Register")
    salary.merge_cells("A1:G1")
    salary["A1"] = "Sahyog Foundation — Consolidated Salary Register FY 2025-26"
    salary["A1"].font = Font(bold=True, size=14)
    salary["A2"] = (
        "Use only payroll line IDs beginning PAY-. Paid lines only; cost-centre "
        "codes may have inconsistent case or surrounding spaces."
    )
    salary.merge_cells("A2:G2")
    salary["A3"] = "Confidential — programme costing copy"
    salary.append([])
    salary.append([])
    salary_headers = [
        "Payroll Line ID",
        "Pay Month",
        "Employee ID",
        "Role",
        "Cost Centre",
        "Gross Salary (INR)",
        "Payroll Status",
    ]
    salary.append(salary_headers)
    style_header(salary, 6)
    section_total = 0
    for index, row in enumerate(salary_records, start=1):
        month_value = (
            row["month"].strftime("%b-%Y")
            if index % 3
            else row["month"].strftime("%Y-%m")
        )
        salary.append(
            [
                row["id"],
                month_value,
                row["employee_id"],
                row["role"],
                messy_code(row["code"], index + 5),
                display_amount(row["amount"], index + 3),
                messy_status(row["status"], index + 2),
            ]
        )
        section_total += row["amount"]
        if index % 2_000 == 0 and index != len(salary_records):
            salary.append([f"Salary section subtotal {index // 2000}", None, None, None, None, section_total, None])
            salary.append([])
            salary.append(salary_headers)
            section_total = 0

    lookup = wb.create_sheet("Cost Centre Lookup")
    lookup.merge_cells("A1:E1")
    lookup["A1"] = "Approved Cost-Centre Mapping — FY 2025-26"
    lookup["A1"].font = Font(bold=True, size=14)
    lookup["A2"] = "Match codes after trimming spaces and ignoring case. Unlisted codes are unmapped."
    lookup.merge_cells("A2:E2")
    lookup.append([])
    lookup.append(["Cost Centre Code", "Programme", "District", "State", "Primary Funder"])
    style_header(lookup, 4)
    for index, values in enumerate(LOOKUP, start=1):
        code, programme, district, state, donor = values
        lookup.append([messy_code(code, index + 9), programme, district, state, donor])

    for ws, widths in (
        (field, [20, 18, 24, 24, 30, 18, 18]),
        (salary, [20, 16, 16, 25, 24, 20, 18]),
        (lookup, [24, 30, 20, 16, 32]),
    ):
        for column, width in zip("ABCDEFG", widths):
            ws.column_dimensions[column].width = width
        ws.freeze_panes = "A6" if ws is field else ("A7" if ws is salary else "A5")

    wb.save(FILES / WORKBOOK_NAME)


def make_case(truth: dict) -> dict:
    join_rule = (
        "Normalize cost-centre codes by trimming spaces and matching case-insensitively "
        "to the Cost Centre Lookup; exclude codes absent from the lookup. Include only "
        "Posted field expenses and Paid salary lines, matching statuses after trimming "
        "spaces and ignoring case."
    )
    return {
        "id": "big_ledger",
        "description": (
            "A 20,000-row national NGO ledger split between field expenses and salary "
            "registers, with messy codes, status text, subtotal rows, and a lookup join."
        ),
        "files": [WORKBOOK_NAME],
        "questions": [
            {
                "id": "q1",
                "question": (
                    "What is the combined eligible field-expense and salary total in INR? "
                    + join_rule
                ),
                "answer": truth["q1"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Summing batch subtotals or including draft, void, held, reversed, or unmapped lines.",
            },
            {
                "id": "q2",
                "question": (
                    "What is the combined eligible amount for Maternal and Child Health "
                    "in Odisha? " + join_rule
                ),
                "answer": truth["q2"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Filtering programme names without first joining cost-centre codes.",
            },
            {
                "id": "q3",
                "question": (
                    "How many eligible transaction lines map to Barmer district across "
                    "both registers? Count field and payroll lines, not employees. " + join_rule
                ),
                "answer": truth["q3"],
                "kind": "number",
                "tol_pct": 0,
                "trap": "Counting distinct employees or including subtotal rows.",
            },
            {
                "id": "q4",
                "question": (
                    "Which programme has the highest combined eligible amount across both "
                    "registers? " + join_rule
                ),
                "answer": truth["q4"],
                "kind": "name",
                "tol_pct": 0,
                "trap": "Ranking raw cost-centre codes instead of programmes from the lookup.",
            },
            {
                "id": "q5",
                "question": (
                    "For Rural Livelihoods, what percentage of its combined eligible amount "
                    "comes from the salary register? Return a percentage from 0 to 100. "
                    + join_rule
                ),
                "answer": truth["q5"],
                "kind": "number",
                "tol_pct": 0.001,
                "trap": "Dividing salary spend by the whole ledger instead of Rural Livelihoods spend.",
            },
            {
                "id": "q6",
                "question": (
                    "Which district has the highest combined eligible amount across both "
                    "registers? " + join_rule
                ),
                "answer": truth["q6"],
                "kind": "name",
                "tol_pct": 0,
                "trap": "Using state or programme totals instead of the district mapping.",
            },
        ],
    }


def main() -> None:
    field_records, salary_records = make_clean_records()
    truth = compute_truth(field_records, salary_records)
    case = make_case(truth)
    write_workbook(field_records, salary_records)
    (ROOT / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        f"big_ledger: generated {len(field_records)} field + "
        f"{len(salary_records)} salary rows, {len(case['questions'])} questions"
    )


if __name__ == "__main__":
    main()
