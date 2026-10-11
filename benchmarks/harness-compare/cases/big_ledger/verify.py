#!/usr/bin/env python3
"""Independently verify every big_ledger answer from the messy workbook."""

from __future__ import annotations

import json
import math
import posixpath
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "files" / "national_programme_ledger.xlsx"
FIELD_ID = re.compile(r"^FE-\d{6}$")
SALARY_ID = re.compile(r"^PAY-\d{6}$")


def normalise_code_openpyxl(value) -> str:
    return str(value).strip().upper()


def amount_openpyxl(value) -> int:
    if value is None:
        raise ValueError("transaction amount is blank")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    text = str(value).strip().replace("₹", "").replace(",", "")
    return int(round(float(text)))


def derive_with_openpyxl() -> dict[str, int | float | str]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    lookup = {}
    for row in wb["Cost Centre Lookup"].iter_rows(values_only=True):
        raw_code = row[0] if row else None
        if raw_code is None:
            continue
        code = normalise_code_openpyxl(raw_code)
        if not code.startswith("CC-") or row[1] in (None, "Programme"):
            continue
        lookup[code] = {
            "programme": str(row[1]).strip(),
            "district": str(row[2]).strip(),
            "state": str(row[3]).strip(),
        }
    if len(lookup) != 24:
        raise ValueError(f"openpyxl lookup has {len(lookup)} codes, expected 24")

    joined = []
    field_count = 0
    for row in wb["Field Expenses"].iter_rows(values_only=True):
        record_id = str(row[0]).strip() if row and row[0] is not None else ""
        if not FIELD_ID.fullmatch(record_id):
            continue
        field_count += 1
        code = normalise_code_openpyxl(row[2])
        mapping = lookup.get(code)
        status = str(row[6]).strip().casefold()
        if mapping is not None and status == "posted":
            joined.append(("field", amount_openpyxl(row[5]), mapping))

    salary_count = 0
    for row in wb["Salary Register"].iter_rows(values_only=True):
        record_id = str(row[0]).strip() if row and row[0] is not None else ""
        if not SALARY_ID.fullmatch(record_id):
            continue
        salary_count += 1
        code = normalise_code_openpyxl(row[4])
        mapping = lookup.get(code)
        status = str(row[6]).strip().casefold()
        if mapping is not None and status == "paid":
            joined.append(("salary", amount_openpyxl(row[5]), mapping))
    wb.close()
    if (field_count, salary_count) != (12_000, 8_000):
        raise ValueError(
            f"openpyxl found {field_count} field and {salary_count} salary rows"
        )

    total = sum(amount for _, amount, _ in joined)
    q2 = sum(
        amount
        for _, amount, mapping in joined
        if mapping["programme"] == "Maternal and Child Health"
        and mapping["state"] == "Odisha"
    )
    q3 = sum(mapping["district"] == "Barmer" for _, _, mapping in joined)
    by_programme = defaultdict(int)
    by_district = defaultdict(int)
    rural_total = 0
    rural_salary = 0
    for source, amount, mapping in joined:
        by_programme[mapping["programme"]] += amount
        by_district[mapping["district"]] += amount
        if mapping["programme"] == "Rural Livelihoods":
            rural_total += amount
            if source == "salary":
                rural_salary += amount
    return {
        "q1": total,
        "q2": q2,
        "q3": q3,
        "q4": max(by_programme.items(), key=lambda pair: pair[1])[0],
        "q5": round(100.0 * rural_salary / rural_total, 6),
        "q6": max(by_district.items(), key=lambda pair: pair[1])[0],
    }


def clean_code_ooxml(value) -> str:
    if value is None:
        return ""
    return re.sub(r"^\s+|\s+$", "", str(value)).casefold()


def amount_ooxml(value) -> int:
    if value is None:
        raise ValueError("transaction amount is blank")
    if isinstance(value, (int, float)):
        return int(round(float(value)))
    digits = re.sub(r"[₹,\s]", "", str(value))
    if not re.fullmatch(r"\d+(?:\.\d+)?", digits):
        raise ValueError(f"invalid transaction amount {value!r}")
    return int(round(float(digits)))


def column_index(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference).group(0)
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result - 1


def read_ooxml_rows(sheet_name: str) -> list[list[object]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(WORKBOOK) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relation_id = None
        for sheet in workbook_root.findall(f".//{{{main_ns}}}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                relation_id = sheet.attrib[f"{{{office_rel_ns}}}id"]
                break
        if relation_id is None:
            raise ValueError(f"sheet not found: {sheet_name}")
        relations = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = None
        for relation in relations.findall(f"{{{package_rel_ns}}}Relationship"):
            if relation.attrib.get("Id") == relation_id:
                target = relation.attrib["Target"]
                break
        if target is None:
            raise ValueError(f"relationship not found for sheet: {sheet_name}")
        sheet_path = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))
        sheet_root = ElementTree.fromstring(archive.read(sheet_path))
        output = []
        for xml_row in sheet_root.findall(f".//{{{main_ns}}}row"):
            values: list[object] = [None] * 7
            for cell in xml_row.findall(f"{{{main_ns}}}c"):
                index = column_index(cell.attrib["r"])
                if index >= len(values):
                    continue
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    inline = cell.find(f"{{{main_ns}}}is")
                    value = "".join(
                        node.text or "" for node in inline.iter(f"{{{main_ns}}}t")
                    )
                else:
                    value_node = cell.find(f"{{{main_ns}}}v")
                    if value_node is None:
                        value = None
                    elif cell_type == "s":
                        value = shared[int(value_node.text)]
                    elif cell_type in ("str", "e"):
                        value = value_node.text
                    else:
                        number = float(value_node.text)
                        value = int(number) if number.is_integer() else number
                values[index] = value
            output.append(values)
        return output


def derive_with_ooxml() -> dict[str, int | float | str]:
    lookup: dict[str, tuple[str, str, str]] = {}
    for row in read_ooxml_rows("Cost Centre Lookup"):
        code = clean_code_ooxml(row[0])
        if not code.startswith("cc-") or str(row[1]).strip() == "Programme":
            continue
        lookup[code] = (str(row[1]).strip(), str(row[2]).strip(), str(row[3]).strip())
    if len(lookup) != 24:
        raise ValueError(f"OOXML lookup has {len(lookup)} codes, expected 24")

    eligible: list[tuple[str, int, str, str, str]] = []
    field_count = 0
    for row in read_ooxml_rows("Field Expenses"):
        record_id = "" if row[0] is None else str(row[0]).strip()
        if not FIELD_ID.fullmatch(record_id):
            continue
        field_count += 1
        mapping = lookup.get(clean_code_ooxml(row[2]))
        status = "" if row[6] is None else str(row[6]).strip().lower()
        if mapping and status == "posted":
            eligible.append(("field", amount_ooxml(row[5]), *mapping))

    salary_count = 0
    for row in read_ooxml_rows("Salary Register"):
        record_id = "" if row[0] is None else str(row[0]).strip()
        if not SALARY_ID.fullmatch(record_id):
            continue
        salary_count += 1
        mapping = lookup.get(clean_code_ooxml(row[4]))
        status = "" if row[6] is None else str(row[6]).strip().lower()
        if mapping and status == "paid":
            eligible.append(("salary", amount_ooxml(row[5]), *mapping))
    if field_count != 12_000 or salary_count != 8_000:
        raise ValueError(
            f"OOXML found {field_count} field and {salary_count} salary rows"
        )

    total = sum(row[1] for row in eligible)
    maternal_odisha = sum(
        row[1]
        for row in eligible
        if row[2] == "Maternal and Child Health" and row[4] == "Odisha"
    )
    barmer_count = len([row for row in eligible if row[3] == "Barmer"])
    programme_totals: dict[str, int] = {}
    district_totals: dict[str, int] = {}
    for _, amount, programme, district, _ in eligible:
        programme_totals[programme] = programme_totals.get(programme, 0) + amount
        district_totals[district] = district_totals.get(district, 0) + amount
    highest_programme = sorted(
        programme_totals, key=lambda key: (-programme_totals[key], key)
    )[0]
    rural_rows = [row for row in eligible if row[2] == "Rural Livelihoods"]
    salary_share = round(
        100.0
        * sum(row[1] for row in rural_rows if row[0] == "salary")
        / sum(row[1] for row in rural_rows),
        6,
    )
    highest_district = sorted(
        district_totals, key=lambda key: (-district_totals[key], key)
    )[0]
    return {
        "q1": total,
        "q2": maternal_odisha,
        "q3": barmer_count,
        "q4": highest_programme,
        "q5": salary_share,
        "q6": highest_district,
    }


def values_match(left, right, tol_pct: float = 0.0) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        return left == right
    left_number = float(left)
    right_number = float(right)
    if tol_pct == 0:
        return math.isclose(left_number, right_number, rel_tol=0, abs_tol=1e-9)
    return abs(left_number - right_number) <= max(abs(right_number), 1.0) * tol_pct / 100.0


def main() -> int:
    try:
        case = json.loads((ROOT / "case.json").read_text(encoding="utf-8"))
        first = derive_with_openpyxl()
        second = derive_with_ooxml()
        errors = []

        if set(first) != set(second):
            errors.append("independent derivations produced different question sets")
        for question_id in sorted(set(first) | set(second)):
            if question_id in first and question_id in second:
                if not values_match(first[question_id], second[question_id]):
                    errors.append(
                        f"{question_id}: derivations disagree "
                        f"({first[question_id]!r} vs {second[question_id]!r})"
                    )

        questions = case.get("questions", [])
        if len(questions) not in (5, 6):
            errors.append(f"case has {len(questions)} questions; expected 5 or 6")
        expected_ids = {question["id"] for question in questions}
        if expected_ids != set(first):
            errors.append(
                f"question IDs {sorted(expected_ids)} do not match derived IDs {sorted(first)}"
            )
        for question in questions:
            question_id = question["id"]
            if question_id not in first:
                continue
            if not values_match(
                first[question_id], question["answer"], float(question.get("tol_pct", 0))
            ):
                errors.append(
                    f"{question_id}: derived {first[question_id]!r}, "
                    f"case.json has {question['answer']!r}"
                )

        if errors:
            for error in errors:
                print(f"ERROR: {error}", file=sys.stderr)
            return 1
        print(
            f"big_ledger verify: PASS ({len(questions)} questions; 12,000 field + "
            "8,000 salary rows; openpyxl and OOXML derivations agree)"
        )
        return 0
    except Exception as exc:
        print(f"ERROR: verification failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
