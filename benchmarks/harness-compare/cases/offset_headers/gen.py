#!/usr/bin/env python3
"""Generate the deterministic offset_headers benchmark case."""

from __future__ import annotations

import json
import random
import shutil
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SEED = 20250623
HERE = Path(__file__).resolve().parent
FILES = HERE / "files"
OUTPUT = FILES / "programme_monthly_spend.xlsx"
MONTHS = ["Apr 2025", "May 2025", "Jun 2025", "Jul 2025", "Aug 2025", "Sep 2025"]

PROGRAMMES = [
    ("PRG-01", "Jal Jeevan Saathi", "Maharashtra", "Jal Jeevan Mission", 1.72),
    ("PRG-02", "Poshan Shakti", "Maharashtra", "Tata Trusts CSR", 1.08),
    ("PRG-03", "Shiksha Udaan", "Rajasthan", "Azim Premji Foundation", 0.96),
    ("PRG-04", "Swasthya Sathi", "Odisha", "Global Health Access Fund", 1.22),
    ("PRG-05", "Mahila Ajeevika", "Bihar", "HDFC Bank Parivartan", 0.91),
    ("PRG-06", "Krishi Pragati", "Odisha", "NABARD Watershed Fund", 1.05),
    ("PRG-07", "Bal Suraksha", "Jharkhand", "Tata Trusts CSR", 0.84),
    ("PRG-08", "Digital Sakhi", "Madhya Pradesh", "HDFC Bank Parivartan", 0.79),
    ("PRG-09", "Gram Urja", "Rajasthan", "Global Green Fund", 1.14),
    ("PRG-10", "Adivasi Adhikar", "Maharashtra", "Azim Premji Foundation", 0.87),
    ("PRG-11", "Nayi Disha Skills", "Bihar", "National Skill Development Mission", 0.93),
    ("PRG-12", "Maa aur Shishu", "Jharkhand", "Global Health Access Fund", 1.16),
]


def make_records() -> list[dict]:
    """Create clean programme-month records before introducing workbook mess."""
    random.seed(SEED)
    rng = np.random.default_rng(SEED)
    records: list[dict] = []
    for programme_index, (code, name, state, donor, weight) in enumerate(PROGRAMMES):
        monthly: dict[str, int] = {}
        for month_index, month in enumerate(MONTHS):
            if (programme_index * 3 + month_index * 2) % 11 == 0 and programme_index != 0:
                value = 0
            else:
                seasonal = [0.91, 1.02, 1.10, 0.95, 1.18, 1.07][month_index]
                value = int(
                    round(rng.integers(410_000, 1_080_000) * weight * seasonal / 500)
                    * 500
                )
            monthly[month] = value
        records.append(
            {
                "code": code,
                "programme": name,
                "state": state,
                "donor": donor,
                "monthly": monthly,
            }
        )
    return records


def answers_from_clean(records: list[dict]) -> dict[str, int | float | str]:
    row_total = lambda record: sum(record["monthly"].values())
    total = sum(row_total(record) for record in records)
    q1_months = MONTHS[:3]
    q2_months = MONTHS[3:]
    maharashtra_q1 = sum(
        record["monthly"][month]
        for record in records
        if record["state"] == "Maharashtra"
        for month in q1_months
    )
    august_positive = sum(record["monthly"]["Aug 2025"] > 0 for record in records)
    programme_totals = {record["programme"]: row_total(record) for record in records}
    donor_totals: dict[str, int] = {}
    for record in records:
        donor_totals[record["donor"]] = donor_totals.get(record["donor"], 0) + row_total(record)
    q2_total = sum(record["monthly"][month] for record in records for month in q2_months)
    return {
        "q1": total,
        "q2": maharashtra_q1,
        "q3": august_positive,
        "q4": max(programme_totals.items(), key=lambda item: (item[1], item[0]))[0],
        "q5": round(100.0 * q2_total / total, 2),
        "q6": max(donor_totals.items(), key=lambda item: (item[1], item[0]))[0],
    }


def messy_zero(programme_index: int, month_index: int) -> object:
    return ("-", "–", None)[(programme_index + month_index) % 3]


def write_workbook(records: list[dict]) -> None:
    if FILES.exists():
        shutil.rmtree(FILES)
    FILES.mkdir(parents=True)

    wb = Workbook()
    wb.properties.creator = "JanSahyog Collective"
    wb.properties.created = datetime(2025, 10, 6, 9, 30, 0)
    wb.properties.modified = datetime(2025, 10, 6, 9, 30, 0)
    ws = wb.active
    ws.title = "Monthly Spend"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:K1")
    ws["A1"] = "JanSahyog Collective — Programme Monthly Spend"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1B5E20")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 27

    ws.merge_cells("A2:K2")
    ws["A2"] = "Prepared by: Grants & Finance Team | Internal management statement"
    ws["A2"].font = Font(italic=True, color="555555")

    ws.merge_cells("A3:K3")
    ws["A3"] = "Reporting period: 01-Apr-2025 to 30-Sep-2025 | Prepared on: 06-Oct-2025 | INR; '-' / '–' / blank means zero"
    ws["A3"].font = Font(size=10, color="555555")

    metadata_headers = [
        ("A4:A5", "Programme Code"),
        ("B4:B5", "Programme"),
        ("C4:C5", "Lead State"),
        ("D4:D5", "Funding Partner"),
        ("K4:K5", "Six-month Total (INR)"),
    ]
    for cell_range, label in metadata_headers:
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.value = label
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("E4:G4")
    ws["E4"] = "Q1 FY 2025–26"
    ws.merge_cells("H4:J4")
    ws["H4"] = "Q2 FY 2025–26"
    for column, month in enumerate(MONTHS, 5):
        ws.cell(5, column, month)

    header_fill = PatternFill("solid", fgColor="2E7D32")
    thin_white = Side(style="thin", color="FFFFFF")
    for row in ws.iter_rows(min_row=4, max_row=5, min_col=1, max_col=11):
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin_white, right=thin_white, top=thin_white, bottom=thin_white)
    ws.row_dimensions[4].height = 23
    ws.row_dimensions[5].height = 31

    for programme_index, record in enumerate(records):
        values: list[object] = [
            record["code"],
            record["programme"],
            record["state"],
            record["donor"],
        ]
        for month_index, month in enumerate(MONTHS):
            value = record["monthly"][month]
            values.append(value if value else messy_zero(programme_index, month_index))
        values.append(sum(record["monthly"].values()))
        ws.append(values)
        for column in range(5, 12):
            if isinstance(ws.cell(ws.max_row, column).value, (int, float)):
                ws.cell(ws.max_row, column).number_format = '₹#,##0;[Red]-₹#,##0'
        if programme_index % 2:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="E8F5E9")

    ws.append([])
    total_row = ws.max_row + 1
    portfolio_months = [
        sum(record["monthly"][month] for record in records) for month in MONTHS
    ]
    ws.append(["PORTFOLIO TOTAL", None, None, None, *portfolio_months, sum(portfolio_months)])
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    for cell in ws[total_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B5E20")
    for column in range(5, 12):
        ws.cell(total_row, column).number_format = '₹#,##0;[Red]-₹#,##0'

    widths = [17, 23, 18, 32, 15, 15, 15, 15, 15, 15, 22]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "E6"
    wb.save(OUTPUT)


def write_case_json(answers: dict[str, int | float | str]) -> None:
    questions = [
        {
            "id": "q1",
            "question": "Using only the programme detail rows (not the PORTFOLIO TOTAL row), what was total spend from April through September 2025, in INR? Treat '-', '–', and blanks as zero.",
            "answer": answers["q1"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "The title rows and two header levels offset the table, while including PORTFOLIO TOTAL doubles the result.",
        },
        {
            "id": "q2",
            "question": "For programme rows whose Lead State is Maharashtra, what was total Q1 FY 2025–26 spend (April, May, and June), in INR? Treat '-', '–', and blanks as zero.",
            "answer": answers["q2"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "A flat one-row header import can lose the quarter-to-month relationship or mistake notes for headers.",
        },
        {
            "id": "q3",
            "question": "How many programme detail rows reported spending greater than zero in August 2025? Treat '-', '–', and blanks as zero and exclude PORTFOLIO TOTAL.",
            "answer": answers["q3"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Counting the summary row or nonnumeric zero markers inflates the programme count.",
        },
        {
            "id": "q4",
            "question": "Which programme had the highest total spend across April through September 2025? Use programme detail rows only and treat '-', '–', and blanks as zero.",
            "answer": answers["q4"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "The merged quarter headers can shift month columns if they are read as an ordinary single header row.",
        },
        {
            "id": "q5",
            "question": "What percentage of total April–September spend occurred in Q2 FY 2025–26 (July, August, and September)? Use detail rows only, treat zero markers as zero, and report to two decimal places.",
            "answer": answers["q5"],
            "kind": "number",
            "tol_pct": 0.01,
            "trap": "Using the displayed row totals together with month values doubles the denominator.",
        },
        {
            "id": "q6",
            "question": "Which Funding Partner supported the highest combined April–September spend across its programme detail rows? Treat zero markers as zero and exclude PORTFOLIO TOTAL.",
            "answer": answers["q6"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "Reading from row 1 or grouping the portfolio summary as a programme obscures the funding-partner totals.",
        },
    ]
    payload = {
        "id": "offset_headers",
        "description": "A wide programme-by-month workbook with three note rows, merged two-level quarter headers, mixed zero markers, and a portfolio summary.",
        "files": [OUTPUT.name],
        "questions": questions,
    }
    (HERE / "case.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main() -> None:
    records = make_records()
    answers = answers_from_clean(records)
    write_workbook(records)
    write_case_json(answers)
    print(f"offset_headers: generated {len(records)} programme rows and {len(answers)} questions")


if __name__ == "__main__":
    main()
