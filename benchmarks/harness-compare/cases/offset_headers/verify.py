#!/usr/bin/env python3
"""Independently verify every answer in the offset_headers case."""

from __future__ import annotations

import json
import math
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
CASE_PATH = HERE / "case.json"
MONTHS = ["Apr 2025", "May 2025", "Jun 2025", "Jul 2025", "Aug 2025", "Sep 2025"]
PROGRAMME_CODE = re.compile(r"^PRG-\d{2}$")
ZERO_MARKERS = {"", "-", "–", "—"}


def fail(message: str) -> None:
    raise AssertionError(message)


def parse_amount(value: object) -> int:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return 0
    if isinstance(value, str) and value.strip() in ZERO_MARKERS:
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        fail(f"unexpected amount value: {value!r}")
    if not float(value).is_integer():
        fail(f"amount is not whole INR: {value!r}")
    return int(value)


def read_sheet_xml(workbook: Path) -> list[list[object]]:
    """Read sheet1 directly from XLSX XML without openpyxl's cell model."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(workbook) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))

    cells: dict[tuple[int, int], object] = {}
    max_row = max_col = 0
    for cell in root.iter(f"{{{main_ns}}}c"):
        reference = cell.attrib["r"]
        match = re.fullmatch(r"([A-Z]+)(\d+)", reference)
        if not match:
            fail(f"invalid cell reference in worksheet XML: {reference!r}")
        letters, row_text = match.groups()
        column = 0
        for character in letters:
            column = column * 26 + ord(character) - ord("A") + 1
        row = int(row_text)
        cell_type = cell.attrib.get("t", "n")
        if cell_type == "inlineStr":
            inline = cell.find(f"{{{main_ns}}}is")
            value: object = "" if inline is None else "".join(
                node.text or "" for node in inline.iter(f"{{{main_ns}}}t")
            )
        else:
            value_node = cell.find(f"{{{main_ns}}}v")
            raw = None if value_node is None else value_node.text
            if raw is None:
                value = None
            elif cell_type == "s":
                value = shared[int(raw)]
            elif cell_type == "b":
                value = raw == "1"
            elif cell_type in ("str", "e"):
                value = raw
            else:
                number = float(raw)
                value = int(number) if number.is_integer() else number
        cells[(row, column)] = value
        max_row = max(max_row, row)
        max_col = max(max_col, column)
    return [
        [cells.get((row, column)) for column in range(1, max_col + 1)]
        for row in range(1, max_row + 1)
    ]


def validate_record(record: dict) -> dict:
    if not PROGRAMME_CODE.fullmatch(record["code"]):
        fail(f"invalid programme code: {record['code']!r}")
    for field in ("programme", "state", "donor"):
        if not record[field] or record[field] == "None":
            fail(f"{record['code']}: missing {field}")
    if set(record["monthly"]) != set(MONTHS):
        fail(f"{record['code']}: month columns are incomplete")
    return record


def parse_with_openpyxl(workbook: Path) -> list[dict]:
    """Derivation A: find header labels cell-by-cell and select PRG rows."""
    wb = load_workbook(workbook, data_only=False, read_only=False)
    if wb.sheetnames != ["Monthly Spend"]:
        fail(f"unexpected sheets: {wb.sheetnames}")
    ws = wb["Monthly Spend"]
    values = [list(row) for row in ws.iter_rows(values_only=True)]

    month_header_rows = [
        index for index, row in enumerate(values) if set(MONTHS).issubset(set(row))
    ]
    if len(month_header_rows) != 1:
        fail(f"expected one month-header row, found {len(month_header_rows)}")
    month_row = month_header_rows[0]
    if month_row == 0:
        fail("month header has no parent quarter row")
    top_row = month_row - 1
    top_labels = values[top_row]
    month_labels = values[month_row]

    metadata = {}
    for label in ("Programme Code", "Programme", "Lead State", "Funding Partner", "Six-month Total (INR)"):
        positions = [index for index, value in enumerate(top_labels) if value == label]
        if len(positions) != 1:
            fail(f"could not uniquely locate {label!r}")
        metadata[label] = positions[0]
    month_columns = {}
    for month in MONTHS:
        positions = [index for index, value in enumerate(month_labels) if value == month]
        if len(positions) != 1:
            fail(f"could not uniquely locate {month!r}")
        month_columns[month] = positions[0]

    records: list[dict] = []
    for row in values[month_row + 1 :]:
        raw_code = row[metadata["Programme Code"]]
        if not isinstance(raw_code, str) or not PROGRAMME_CODE.fullmatch(raw_code.strip()):
            continue
        record = {
            "code": raw_code.strip(),
            "programme": str(row[metadata["Programme"]]).strip(),
            "state": str(row[metadata["Lead State"]]).strip(),
            "donor": str(row[metadata["Funding Partner"]]).strip(),
            "monthly": {month: parse_amount(row[column]) for month, column in month_columns.items()},
        }
        validate_record(record)
        claimed_total = parse_amount(row[metadata["Six-month Total (INR)"]])
        calculated_total = sum(record["monthly"].values())
        if claimed_total != calculated_total:
            fail(f"{record['code']}: row total {claimed_total} != monthly sum {calculated_total}")
        records.append(record)

    if len({record["code"] for record in records}) != len(records) or not records:
        fail("programme codes are missing or duplicated")

    portfolio_rows = [
        row
        for row in values
        if isinstance(row[metadata["Programme Code"]], str)
        and row[metadata["Programme Code"]].strip() == "PORTFOLIO TOTAL"
    ]
    if len(portfolio_rows) != 1:
        fail(f"expected one portfolio total row, found {len(portfolio_rows)}")
    portfolio = portfolio_rows[0]
    for month, column in month_columns.items():
        claimed = parse_amount(portfolio[column])
        derived = sum(record["monthly"][month] for record in records)
        if claimed != derived:
            fail(f"portfolio {month} total {claimed} != detail sum {derived}")
    claimed_grand = parse_amount(portfolio[metadata["Six-month Total (INR)"]])
    derived_grand = sum(sum(record["monthly"].values()) for record in records)
    if claimed_grand != derived_grand:
        fail(f"portfolio grand total {claimed_grand} != detail sum {derived_grand}")
    return records


def parse_with_raw_xml(workbook: Path) -> list[dict]:
    """Derivation B: parse raw XLSX XML and reconstruct the offset header."""
    rows = read_sheet_xml(workbook)
    month_header_rows = [
        index for index, row in enumerate(rows) if set(MONTHS).issubset(set(row))
    ]
    if len(month_header_rows) != 1 or month_header_rows[0] == 0:
        fail(f"raw XML parser found {len(month_header_rows)} month-header rows")
    month_row = month_header_rows[0]
    top_labels = rows[month_row - 1]
    month_labels = rows[month_row]
    metadata = {}
    for label in ("Programme Code", "Programme", "Lead State", "Funding Partner", "Six-month Total (INR)"):
        positions = [index for index, value in enumerate(top_labels) if value == label]
        if len(positions) != 1:
            fail(f"raw XML parser could not uniquely locate {label!r}")
        metadata[label] = positions[0]
    month_columns = {}
    for month in MONTHS:
        positions = [index for index, value in enumerate(month_labels) if value == month]
        if len(positions) != 1:
            fail(f"raw XML parser could not uniquely locate {month!r}")
        month_columns[month] = positions[0]

    records: list[dict] = []
    for row in rows[month_row + 1 :]:
        raw_code = row[metadata["Programme Code"]]
        code = str(raw_code).strip()
        if not PROGRAMME_CODE.fullmatch(code):
            continue
        record = {
            "code": code,
            "programme": str(row[metadata["Programme"]]).strip(),
            "state": str(row[metadata["Lead State"]]).strip(),
            "donor": str(row[metadata["Funding Partner"]]).strip(),
            "monthly": {month: parse_amount(row[column]) for month, column in month_columns.items()},
        }
        validate_record(record)
        claimed_total = parse_amount(row[metadata["Six-month Total (INR)"]])
        if claimed_total != sum(record["monthly"].values()):
            fail(f"raw XML parser: {code} row total disagrees with months")
        records.append(record)
    if not records:
        fail("raw XML parser found no programme detail rows")
    return records


def derive_with_loops(records: list[dict]) -> dict[str, int | float | str]:
    totals_by_programme: dict[str, int] = {}
    totals_by_donor: dict[str, int] = defaultdict(int)
    grand_total = 0
    q2_total = 0
    for record in records:
        programme_total = sum(record["monthly"].values())
        totals_by_programme[record["programme"]] = programme_total
        totals_by_donor[record["donor"]] += programme_total
        grand_total += programme_total
        q2_total += sum(record["monthly"][month] for month in MONTHS[3:])
    return {
        "q1": grand_total,
        "q2": sum(
            record["monthly"][month]
            for record in records
            if record["state"] == "Maharashtra"
            for month in MONTHS[:3]
        ),
        "q3": sum(record["monthly"]["Aug 2025"] > 0 for record in records),
        "q4": max(totals_by_programme.items(), key=lambda item: (item[1], item[0]))[0],
        "q5": round(100.0 * q2_total / grand_total, 2),
        "q6": max(totals_by_donor.items(), key=lambda item: (item[1], item[0]))[0],
    }


def derive_with_flat_columns(records: list[dict]) -> dict[str, int | float | str]:
    flat_rows = [
        {
            "code": record["code"],
            "programme": record["programme"],
            "state": record["state"],
            "donor": record["donor"],
            **record["monthly"],
        }
        for record in records
    ]
    for row in flat_rows:
        row["calculated_total"] = sum(row[month] for month in MONTHS)
    grand_total = sum(row["calculated_total"] for row in flat_rows)
    programme_totals = {
        name: sum(row["calculated_total"] for row in flat_rows if row["programme"] == name)
        for name in {row["programme"] for row in flat_rows}
    }
    donor_totals = {
        donor: sum(row["calculated_total"] for row in flat_rows if row["donor"] == donor)
        for donor in {row["donor"] for row in flat_rows}
    }
    return {
        "q1": grand_total,
        "q2": sum(
            row[month]
            for row in flat_rows
            if row["state"] == "Maharashtra"
            for month in MONTHS[:3]
        ),
        "q3": len([row for row in flat_rows if row["Aug 2025"] > 0]),
        "q4": max(programme_totals.items(), key=lambda item: (item[1], item[0]))[0],
        "q5": round(
            100.0
            * sum(row[month] for row in flat_rows for month in MONTHS[3:])
            / grand_total,
            2,
        ),
        "q6": max(donor_totals.items(), key=lambda item: (item[1], item[0]))[0],
    }


def values_agree(left: object, right: object, kind: str, tol_pct: float) -> bool:
    if kind == "name":
        return isinstance(left, str) and isinstance(right, str) and left == right
    if isinstance(left, bool) or isinstance(right, bool):
        return False
    try:
        left_number = float(left)
        right_number = float(right)
    except (TypeError, ValueError):
        return False
    tolerance = abs(right_number) * tol_pct / 100.0
    return abs(left_number - right_number) <= max(tolerance, 1e-9)


def main() -> None:
    case = json.loads(CASE_PATH.read_text(encoding="utf-8"))
    if case.get("id") != "offset_headers":
        fail("case id is not offset_headers")
    questions = case.get("questions")
    if not isinstance(questions, list) or len(questions) not in (5, 6):
        fail("case must contain 5 or 6 questions")
    if len({question.get("id") for question in questions}) != len(questions):
        fail("question IDs are not unique")
    files = case.get("files")
    if files != ["programme_monthly_spend.xlsx"]:
        fail(f"unexpected file manifest: {files!r}")
    workbook = HERE / "files" / files[0]
    if not workbook.is_file():
        fail(f"missing workbook: {workbook}")
    if workbook.stat().st_size >= 2_000_000:
        fail(f"workbook is not under 2 MB: {workbook.stat().st_size} bytes")

    records_a = parse_with_openpyxl(workbook)
    records_b = parse_with_raw_xml(workbook)
    canonical = lambda records: sorted(
        (
            record["code"],
            record["programme"],
            record["state"],
            record["donor"],
            tuple((month, record["monthly"][month]) for month in MONTHS),
        )
        for record in records
    )
    if canonical(records_a) != canonical(records_b):
        fail("independent parsers did not recover the same programme-month records")
    answers_a = derive_with_loops(records_a)
    answers_b = derive_with_flat_columns(records_b)

    if set(answers_a) != {question["id"] for question in questions}:
        fail("derived answer IDs do not exactly match case questions")
    for question in questions:
        question_id = question["id"]
        kind = question.get("kind")
        tolerance = float(question.get("tol_pct", 0.0))
        if kind not in ("number", "name"):
            fail(f"{question_id}: invalid kind {kind!r}")
        if not question.get("question") or not question.get("trap"):
            fail(f"{question_id}: missing question or trap")
        if not values_agree(answers_a[question_id], answers_b[question_id], kind, 0.0):
            fail(
                f"{question_id}: independent derivations disagree: "
                f"{answers_a[question_id]!r} != {answers_b[question_id]!r}"
            )
        if not values_agree(answers_a[question_id], question.get("answer"), kind, tolerance):
            fail(
                f"{question_id}: stored answer {question.get('answer')!r} "
                f"does not match derived {answers_a[question_id]!r}"
            )

    print(
        f"offset_headers: PASS — {len(records_a)} programme rows, "
        f"{len(questions)} questions, 2 independent derivations agree"
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"offset_headers: FAIL — {exc}", file=sys.stderr)
        raise SystemExit(1)
