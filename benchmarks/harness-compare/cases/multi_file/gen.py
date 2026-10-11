#!/usr/bin/env python3
"""Generate the deterministic multi_file benchmark case."""

from __future__ import annotations

import csv
import json
import random
import shutil
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SEED = 940_217
BASE_DIR = Path(__file__).resolve().parent
FILES_DIR = BASE_DIR / "files"

LOCATIONS = {
    "Odisha": {
        "Koraput": ["Lamtaput", "Semiliguda"],
        "Rayagada": ["Muniguda", "Gunupur"],
    },
    "Maharashtra": {
        "Gadchiroli": ["Dhanora", "Etapalli"],
        "Nandurbar": ["Akkalkuwa", "Taloda"],
    },
    "Karnataka": {
        "Raichur": ["Devadurga", "Manvi"],
        "Koppal": ["Gangavathi", "Yelburga"],
    },
    "Jharkhand": {
        "Dumka": ["Jama", "Shikaripara"],
        "Khunti": ["Murhu", "Torpa"],
    },
}
PROGRAMMES = [
    "Livelihoods",
    "Health & Nutrition",
    "Girls' Education",
    "Water Security",
]
DONORS = [
    "Tata Trusts",
    "HDFC Bank Parivartan",
    "National Rural Livelihoods Mission",
    "Azim Premji Foundation",
    "GlobalGiving Foreign Grant",
]
FIRST_NAMES = [
    "Aarav",
    "Asha",
    "Deepa",
    "Farida",
    "Geeta",
    "Imran",
    "Jyoti",
    "Kiran",
    "Lakshmi",
    "Manoj",
    "Nirmala",
    "Pooja",
    "Rahul",
    "Salma",
    "Sunita",
    "Vijay",
]
LAST_NAMES = [
    "Barla",
    "Das",
    "Gavade",
    "Hembram",
    "Jadhav",
    "Kisku",
    "Majhi",
    "Naik",
    "Oraon",
    "Patil",
    "Sahu",
    "Toppo",
]


def build_clean_beneficiaries() -> list[dict]:
    rng = random.Random(SEED)
    start = date(2024, 7, 1)
    beneficiaries: list[dict] = []
    for index in range(760):
        state = rng.choices(
            list(LOCATIONS), weights=[0.29, 0.26, 0.24, 0.21], k=1
        )[0]
        district = rng.choice(list(LOCATIONS[state]))
        village = rng.choice(LOCATIONS[state][district])
        beneficiaries.append(
            {
                "beneficiary_id": f"BEN-{index + 1:05d}",
                "full_name": f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}",
                "gender": rng.choices(
                    ["Woman", "Man", "Non-binary"],
                    weights=[0.62, 0.36, 0.02],
                    k=1,
                )[0],
                "state": state,
                "district": district,
                "village": village,
                "programme": rng.choices(
                    PROGRAMMES, weights=[0.34, 0.26, 0.23, 0.17], k=1
                )[0],
                "enrolment_date": start + timedelta(days=rng.randint(0, 364)),
            }
        )
    return beneficiaries


def build_clean_payments(beneficiaries: list[dict]) -> list[dict]:
    rng = random.Random(SEED + 1)
    payment_date = date(2025, 4, 1)
    payments: list[dict] = []
    serial = 1
    amounts = [2_500, 4_000, 5_000, 7_500, 10_000, 12_500, 15_000, 20_000]

    for beneficiary in beneficiaries:
        payment_count = rng.choices([0, 1, 2, 3, 4], [0.14, 0.25, 0.30, 0.21, 0.10])[0]
        for tranche in range(1, payment_count + 1):
            payments.append(
                {
                    "payment_id": f"PMT-2025-{serial:05d}",
                    "beneficiary_id": beneficiary["beneficiary_id"],
                    "payment_date": payment_date
                    + timedelta(days=rng.randint(0, 182)),
                    "tranche": tranche,
                    "donor": rng.choice(DONORS),
                    "amount_inr": rng.choice(amounts),
                    "status": "Paid" if rng.random() < 0.87 else "Failed",
                }
            )
            serial += 1

    # These IDs intentionally have no row in beneficiaries.csv.
    for orphan_index in range(92):
        payments.append(
            {
                "payment_id": f"PMT-2025-{serial:05d}",
                "beneficiary_id": f"ORPH-{orphan_index + 1:04d}",
                "payment_date": payment_date + timedelta(days=rng.randint(0, 182)),
                "tranche": rng.randint(1, 3),
                "donor": rng.choice(DONORS),
                "amount_inr": rng.choice(amounts),
                "status": "Paid" if rng.random() < 0.84 else "Failed",
            }
        )
        serial += 1
    return payments


def derive_clean_answers(
    beneficiaries: list[dict], payments: list[dict]
) -> dict[str, int | float | str]:
    """Derive truth from clean in-memory records before any messy files exist."""
    beneficiary_by_id = {
        beneficiary["beneficiary_id"]: beneficiary for beneficiary in beneficiaries
    }
    joined_paid: list[tuple[dict, dict]] = []
    for payment in payments:
        beneficiary = beneficiary_by_id.get(payment["beneficiary_id"])
        if beneficiary is not None and payment["status"] == "Paid":
            joined_paid.append((payment, beneficiary))

    programme_totals: defaultdict[str, int] = defaultdict(int)
    for payment, beneficiary in joined_paid:
        programme_totals[beneficiary["programme"]] += payment["amount_inr"]
    maximum = max(programme_totals.values())
    winners = [
        programme
        for programme, amount in programme_totals.items()
        if amount == maximum
    ]
    if len(winners) != 1:
        raise RuntimeError(f"Programme ranking is tied: {winners}")

    recipients = {
        payment["beneficiary_id"] for payment, _beneficiary in joined_paid
    }
    answers: dict[str, int | float | str] = {
        "q1": sum(payment["amount_inr"] for payment, _ in joined_paid),
        "q2": sum(
            payment["amount_inr"]
            for payment, beneficiary in joined_paid
            if beneficiary["state"] == "Odisha"
            and beneficiary["programme"] == "Girls' Education"
        ),
        "q3": len(recipients),
        "q4": winners[0],
        "q5": round(100.0 * len(recipients) / len(beneficiaries), 2),
        "q6": sum(
            payment["amount_inr"]
            for payment, beneficiary in joined_paid
            if payment["donor"] == "Azim Premji Foundation"
            and beneficiary["gender"] == "Woman"
        ),
    }
    if not answers["q2"] or not answers["q6"]:
        raise RuntimeError("Filtered payment questions must have non-zero answers")
    return answers


def write_beneficiaries_csv(beneficiaries: list[dict], output_path: Path) -> None:
    fieldnames = [
        "Beneficiary ID",
        "Full Name",
        "Gender",
        "State",
        "District",
        "Village",
        "Programme",
        "Enrolment Date",
    ]
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for beneficiary in beneficiaries:
            writer.writerow(
                {
                    "Beneficiary ID": beneficiary["beneficiary_id"],
                    "Full Name": beneficiary["full_name"],
                    "Gender": beneficiary["gender"],
                    "State": beneficiary["state"],
                    "District": beneficiary["district"],
                    "Village": beneficiary["village"],
                    "Programme": beneficiary["programme"],
                    "Enrolment Date": beneficiary["enrolment_date"].isoformat(),
                }
            )


def display_amount(amount: int, row_number: int) -> int | str:
    if row_number % 8 == 0:
        return f"₹{amount:,}"
    if row_number % 13 == 0:
        return f" {amount:,} "
    return amount


def display_status(status: str, row_number: int) -> str:
    if row_number % 9 == 0:
        return f" {status.upper()} "
    if row_number % 5 == 0:
        return status.lower()
    return status


def write_payments_workbook(payments: list[dict], output_path: Path) -> None:
    workbook = Workbook()
    workbook.properties.creator = "Jan Kalyan Programme Finance"
    workbook.properties.created = datetime(2025, 10, 2, 10, 30, 0)
    workbook.properties.modified = datetime(2025, 10, 2, 10, 30, 0)
    sheet = workbook.active
    sheet.title = "Disbursements"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Jan Kalyan Programme — Beneficiary Disbursement Register"
    sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="7030A0")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:G2")
    sheet["A2"] = (
        "Join Beneficiary ID to beneficiaries.csv; unmatched legacy IDs remain "
        "for finance reconciliation."
    )
    sheet["A2"].font = Font(italic=True, color="666666")
    sheet.merge_cells("A3:G3")
    sheet["A3"] = "FY 2025–26 | Status and amount cells retain field-office formatting"

    headers = [
        "Payment ID",
        "Beneficiary ID",
        "Payment Date",
        "Tranche",
        "Donor",
        "Amount (INR)",
        "Status",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_number = 6
    data_row_number = 0
    for payment in payments:
        data_row_number += 1
        values = [
            payment["payment_id"],
            payment["beneficiary_id"],
            payment["payment_date"],
            payment["tranche"],
            payment["donor"],
            display_amount(payment["amount_inr"], data_row_number),
            display_status(payment["status"], data_row_number),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=3).number_format = "dd-mmm-yyyy"
        row_number += 1
        if data_row_number % 275 == 0:
            row_number += 1

    sheet.cell(row=row_number + 1, column=1, value="END OF REGISTER")
    widths = [20, 19, 15, 11, 35, 18, 13]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:G{row_number - 1}"
    workbook.save(output_path)


def build_questions(answers: dict[str, int | float | str]) -> list[dict]:
    common = (
        "Join payments.xlsx to beneficiaries.csv on Beneficiary ID; include "
        "only status Paid after trimming whitespace and ignoring case, and "
        "exclude payment rows whose Beneficiary ID is absent from the CSV. "
    )
    return [
        {
            "id": "q1",
            "question": common
            + "Treat numeric and ₹/comma-formatted amounts as INR. What is the total paid amount?",
            "answer": answers["q1"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Orphan and failed payments inflate the total if not excluded.",
        },
        {
            "id": "q2",
            "question": common
            + "What total amount was paid to Girls' Education beneficiaries in Odisha?",
            "answer": answers["q2"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Programme and state live in the CSV, so filtering payments alone cannot answer this.",
        },
        {
            "id": "q3",
            "question": common
            + "How many distinct beneficiaries received at least one such paid payment?",
            "answer": answers["q3"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Counting payment rows counts multi-tranche beneficiaries more than once.",
        },
        {
            "id": "q4",
            "question": common
            + "Which programme received the highest total paid amount?",
            "answer": answers["q4"],
            "kind": "name",
            "tol_pct": 0.0,
            "trap": "Ranking requires the beneficiary join and exclusion of orphan and failed rows.",
        },
        {
            "id": "q5",
            "question": common
            + "What percentage of all beneficiaries listed in the CSV received at least one such paid payment? Round to two decimal places.",
            "answer": answers["q5"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "The denominator is all CSV beneficiaries, not payment rows or only recipients.",
        },
        {
            "id": "q6",
            "question": common
            + "What total amount did Azim Premji Foundation pay to beneficiaries whose CSV gender is Woman?",
            "answer": answers["q6"],
            "kind": "number",
            "tol_pct": 0.0,
            "trap": "Donor is in the workbook while gender is in the CSV; orphan and failed rows still must be removed.",
        },
    ]


def main() -> None:
    clean_beneficiaries = build_clean_beneficiaries()
    clean_payments = build_clean_payments(clean_beneficiaries)
    answers = derive_clean_answers(clean_beneficiaries, clean_payments)

    if FILES_DIR.exists():
        shutil.rmtree(FILES_DIR)
    FILES_DIR.mkdir(parents=True)
    beneficiaries_path = FILES_DIR / "beneficiaries.csv"
    payments_path = FILES_DIR / "payments.xlsx"
    write_beneficiaries_csv(clean_beneficiaries, beneficiaries_path)
    write_payments_workbook(clean_payments, payments_path)

    case = {
        "id": "multi_file",
        "description": (
            "A beneficiary roster CSV and a separately formatted payment "
            "workbook requiring an ID join, paid-status filtering and orphan exclusion."
        ),
        "files": [beneficiaries_path.name, payments_path.name],
        "questions": build_questions(answers),
    }
    (BASE_DIR / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    for path in (beneficiaries_path, payments_path):
        if path.stat().st_size >= 2 * 1024 * 1024:
            raise RuntimeError(f"{path.name} exceeds 2 MB")
    print(
        "multi_file: generated beneficiaries.csv "
        f"({beneficiaries_path.stat().st_size} bytes) and payments.xlsx "
        f"({payments_path.stat().st_size} bytes), {len(case['questions'])} questions"
    )


if __name__ == "__main__":
    main()
