#!/usr/bin/env python3
"""Independently verify every multi_file answer from the generated files."""

from __future__ import annotations

import csv
import json
import math
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
EXPECTED_IDS = {"q1", "q2", "q3", "q4", "q5", "q6"}


def derive_with_stdlib_openpyxl(
    beneficiaries_path: Path, payments_path: Path
) -> dict[str, int | float | str]:
    """First derivation: csv.DictReader, openpyxl cells and Python reductions."""
    beneficiaries: dict[str, dict[str, str]] = {}
    with beneficiaries_path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        required = {
            "Beneficiary ID",
            "Gender",
            "State",
            "Programme",
        }
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Missing beneficiary columns: {sorted(missing)}")
        for row in reader:
            beneficiary_id = row["Beneficiary ID"]
            if beneficiary_id in beneficiaries:
                raise ValueError(f"Duplicate beneficiary ID: {beneficiary_id}")
            beneficiaries[beneficiary_id] = row
    if not beneficiaries:
        raise ValueError("No beneficiaries parsed")

    workbook = load_workbook(payments_path, read_only=True, data_only=False)
    sheet = workbook["Disbursements"]
    header_row = None
    headers: dict[str, int] = {}
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        labels = [str(value).strip() if value is not None else "" for value in row]
        if "Payment ID" in labels:
            header_row = row_number
            headers = {label: index for index, label in enumerate(labels) if label}
            break
    if header_row is None:
        raise ValueError("Could not locate payment header")

    required_payment_columns = {
        "Payment ID",
        "Beneficiary ID",
        "Donor",
        "Amount (INR)",
        "Status",
    }
    missing = required_payment_columns - headers.keys()
    if missing:
        raise ValueError(f"Missing payment columns: {sorted(missing)}")

    joined_paid: list[tuple[str, int, str, dict[str, str]]] = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_payment_id = row[headers["Payment ID"]]
        payment_id = "" if raw_payment_id is None else str(raw_payment_id).strip()
        if not payment_id.startswith("PMT-"):
            continue
        beneficiary_id = str(row[headers["Beneficiary ID"]])
        beneficiary = beneficiaries.get(beneficiary_id)
        status = str(row[headers["Status"]]).strip().casefold()
        if beneficiary is None or status != "paid":
            continue

        raw_amount = row[headers["Amount (INR)"]]
        if isinstance(raw_amount, (int, float)) and not isinstance(raw_amount, bool):
            amount = int(raw_amount)
        else:
            amount = int(re.sub(r"[₹,\s]", "", str(raw_amount)))
        donor = str(row[headers["Donor"]]).strip()
        joined_paid.append((beneficiary_id, amount, donor, beneficiary))
    workbook.close()

    programme_totals: defaultdict[str, int] = defaultdict(int)
    for _beneficiary_id, amount, _donor, beneficiary in joined_paid:
        programme_totals[beneficiary["Programme"]] += amount
    maximum = max(programme_totals.values())
    winners = [
        programme
        for programme, amount in programme_totals.items()
        if amount == maximum
    ]
    if len(winners) != 1:
        raise ValueError(f"Programme ranking is not unique: {winners}")

    recipients = {item[0] for item in joined_paid}
    return {
        "q1": sum(item[1] for item in joined_paid),
        "q2": sum(
            amount
            for _beneficiary_id, amount, _donor, beneficiary in joined_paid
            if beneficiary["State"] == "Odisha"
            and beneficiary["Programme"] == "Girls' Education"
        ),
        "q3": len(recipients),
        "q4": winners[0],
        "q5": round(100.0 * len(recipients) / len(beneficiaries), 2),
        "q6": sum(
            amount
            for _beneficiary_id, amount, donor, beneficiary in joined_paid
            if donor == "Azim Premji Foundation"
            and beneficiary["Gender"] == "Woman"
        ),
    }


def raw_xlsx_rows(path: Path) -> list[dict[int, object]]:
    """Parse the first worksheet directly from XLSX XML, without openpyxl."""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    qualified = lambda name: f"{{{namespace}}}{name}"
    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(qualified("si")):
                shared_strings.append(
                    "".join(
                        node.text or "" for node in item.iter(qualified("t"))
                    )
                )
        sheet_root = ElementTree.fromstring(
            archive.read("xl/worksheets/sheet1.xml")
        )

    parsed_rows: list[dict[int, object]] = []
    for xml_row in sheet_root.findall(
        f".//{qualified('sheetData')}/{qualified('row')}"
    ):
        parsed: dict[int, object] = {}
        for cell in xml_row.findall(qualified("c")):
            reference = cell.attrib["r"]
            letters = re.match(r"[A-Z]+", reference)
            if letters is None:
                raise ValueError(f"Bad XLSX cell reference: {reference}")
            column = 0
            for letter in letters.group(0):
                column = column * 26 + ord(letter) - ord("A") + 1

            cell_type = cell.attrib.get("t", "n")
            if cell_type == "inlineStr":
                value: object = "".join(
                    node.text or "" for node in cell.iter(qualified("t"))
                )
            else:
                value_node = cell.find(qualified("v"))
                raw = "" if value_node is None else value_node.text or ""
                if cell_type == "s":
                    value = shared_strings[int(raw)]
                elif cell_type in {"str", "e"}:
                    value = raw
                elif cell_type == "b":
                    value = raw == "1"
                elif raw == "":
                    value = None
                else:
                    number = float(raw)
                    value = int(number) if number.is_integer() else number
            parsed[column] = value
        parsed_rows.append(parsed)
    return parsed_rows


def derive_with_csv_reader_and_xml(
    beneficiaries_path: Path, payments_path: Path
) -> dict[str, int | float | str]:
    """Second derivation: positional CSV parsing, raw XLSX XML and accumulators."""
    with beneficiaries_path.open(newline="", encoding="utf-8") as handle:
        csv_rows = list(csv.reader(handle))
    if not csv_rows:
        raise ValueError("Beneficiary CSV is empty")
    csv_columns = {name: index for index, name in enumerate(csv_rows[0])}
    beneficiaries: dict[str, tuple[str, str, str]] = {}
    for row in csv_rows[1:]:
        beneficiary_id = row[csv_columns["Beneficiary ID"]]
        if beneficiary_id in beneficiaries:
            raise ValueError(f"Duplicate beneficiary ID: {beneficiary_id}")
        beneficiaries[beneficiary_id] = (
            row[csv_columns["State"]],
            row[csv_columns["Programme"]],
            row[csv_columns["Gender"]],
        )

    xlsx_rows = raw_xlsx_rows(payments_path)
    header_position = None
    columns: dict[str, int] = {}
    for position, row in enumerate(xlsx_rows):
        labels = {
            str(value).strip(): column
            for column, value in row.items()
            if value is not None
        }
        if "Payment ID" in labels:
            header_position = position
            columns = labels
            break
    if header_position is None:
        raise ValueError("XML parser could not locate Payment ID header")

    programme_totals: defaultdict[str, int] = defaultdict(int)
    recipients: set[str] = set()
    total_paid = odisha_girls_paid = azim_women_paid = 0
    for row in xlsx_rows[header_position + 1 :]:
        payment_id = str(row.get(columns["Payment ID"], "")).strip()
        if not payment_id.startswith("PMT-"):
            continue
        beneficiary_id = str(row[columns["Beneficiary ID"]])
        beneficiary = beneficiaries.get(beneficiary_id)
        status = str(row[columns["Status"]]).strip().casefold()
        if beneficiary is None or status != "paid":
            continue

        raw_amount = row[columns["Amount (INR)"]]
        amount = (
            int(raw_amount)
            if isinstance(raw_amount, (int, float)) and not isinstance(raw_amount, bool)
            else int(re.sub(r"[₹,\s]", "", str(raw_amount)))
        )
        state, programme, gender = beneficiary
        donor = str(row[columns["Donor"]]).strip()
        recipients.add(beneficiary_id)
        total_paid += amount
        programme_totals[programme] += amount
        if state == "Odisha" and programme == "Girls' Education":
            odisha_girls_paid += amount
        if donor == "Azim Premji Foundation" and gender == "Woman":
            azim_women_paid += amount

    maximum = max(programme_totals.values())
    winners = [
        programme
        for programme, amount in programme_totals.items()
        if amount == maximum
    ]
    if len(winners) != 1:
        raise ValueError(f"Programme ranking is not unique: {winners}")
    return {
        "q1": total_paid,
        "q2": odisha_girls_paid,
        "q3": len(recipients),
        "q4": winners[0],
        "q5": round(100.0 * len(recipients) / len(beneficiaries), 2),
        "q6": azim_women_paid,
    }


def answers_match(actual: object, expected: object, kind: str, tol_pct: float) -> bool:
    if kind == "name":
        return isinstance(actual, str) and actual == expected
    if not isinstance(actual, (int, float)) or isinstance(actual, bool):
        return False
    if not isinstance(expected, (int, float)) or isinstance(expected, bool):
        return False
    allowed = abs(float(expected)) * tol_pct / 100.0
    return math.isclose(
        float(actual), float(expected), rel_tol=0.0, abs_tol=max(allowed, 1e-9)
    )


def main() -> int:
    try:
        case = json.loads((BASE_DIR / "case.json").read_text(encoding="utf-8"))
        questions = case.get("questions", [])
        if len(questions) not in (5, 6):
            raise ValueError(f"Expected 5 or 6 questions, found {len(questions)}")
        ids = {question.get("id") for question in questions}
        if ids != EXPECTED_IDS:
            raise ValueError(f"Unexpected question IDs: {sorted(ids)}")
        expected_files = ["beneficiaries.csv", "payments.xlsx"]
        if case.get("files") != expected_files:
            raise ValueError(f"Unexpected files declaration: {case.get('files')}")

        beneficiaries_path = BASE_DIR / "files" / "beneficiaries.csv"
        payments_path = BASE_DIR / "files" / "payments.xlsx"
        for path in (beneficiaries_path, payments_path):
            if path.stat().st_size >= 2 * 1024 * 1024:
                raise ValueError(f"{path.name} is 2 MB or larger")

        first = derive_with_stdlib_openpyxl(
            beneficiaries_path, payments_path
        )
        second = derive_with_csv_reader_and_xml(
            beneficiaries_path, payments_path
        )
        errors: list[str] = []
        for question in questions:
            question_id = question["id"]
            if not answers_match(
                first[question_id],
                second[question_id],
                question["kind"],
                0.0,
            ):
                errors.append(
                    f"{question_id}: independent derivations disagree "
                    f"({first[question_id]!r} != {second[question_id]!r})"
                )
            if not answers_match(
                first[question_id],
                question["answer"],
                question["kind"],
                float(question.get("tol_pct", 0.0)),
            ):
                errors.append(
                    f"{question_id}: case answer {question['answer']!r}, "
                    f"parsed answer {first[question_id]!r}"
                )

        if errors:
            for error in errors:
                print(f"ERROR: {error}", file=sys.stderr)
            return 1
        print(
            f"multi_file: PASS — {len(questions)} questions verified by "
            "independent openpyxl/DictReader and raw-XML/csv.reader derivations"
        )
        return 0
    except Exception as exc:
        print(f"multi_file: FAIL — {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
