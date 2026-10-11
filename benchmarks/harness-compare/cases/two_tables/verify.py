#!/usr/bin/env python3
"""Independently verify every two_tables answer from the messy workbook."""

from __future__ import annotations

import json
import math
import posixpath
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
CASE_PATH = ROOT / "case.json"


def parse_with_openpyxl(path: Path) -> tuple[list[dict], list[dict]]:
    ws = load_workbook(path, data_only=False, read_only=True)["Finance FY25-26"]
    rows = list(ws.iter_rows(values_only=True))
    budget_header = next(
        i for i, row in enumerate(rows) if row and row[0] == "Budget Code"
    )
    receipt_header = next(
        i for i, row in enumerate(rows) if row and row[0] == "Receipt ID"
    )

    budgets: list[dict] = []
    for row in rows[budget_header + 1 : receipt_header]:
        if not row or row[0] in (None, "TOTAL APPROVED"):
            continue
        if not str(row[0]).startswith("BG-"):
            continue
        budgets.append(
            {
                "code": str(row[0]),
                "programme": str(row[1]),
                "approved": int(row[4]),
            }
        )

    receipts: list[dict] = []
    for row in rows[receipt_header + 1 :]:
        if not row or row[0] in (None, "TOTAL RECEIPTS"):
            continue
        if not str(row[0]).startswith("RCPT-"):
            continue
        receipts.append(
            {
                "receipt_id": str(row[0]),
                "donor": str(row[2]),
                "donor_type": str(row[3]),
                "programme": str(row[4]),
                "amount": int(row[5]),
            }
        )
    if not budgets or not receipts:
        raise ValueError("Could not locate both detail tables with openpyxl")
    return budgets, receipts


def derive_with_loops(budgets: list[dict], receipts: list[dict]) -> dict:
    total_budget = 0
    for row in budgets:
        total_budget += row["approved"]

    jal_receipts = 0
    foreign_ids: set[str] = set()
    donor_totals: dict[str, int] = {}
    total_receipts = 0
    for row in receipts:
        amount = row["amount"]
        total_receipts += amount
        if row["programme"] == "Jal Jeevan Saathi":
            jal_receipts += amount
        if row["donor_type"] == "Foreign Grant":
            foreign_ids.add(row["receipt_id"])
        donor_totals[row["donor"]] = donor_totals.get(row["donor"], 0) + amount
    ranking = sorted(donor_totals.items(), key=lambda item: (-item[1], item[0]))
    if len(ranking) > 1 and ranking[0][1] == ranking[1][1]:
        raise ValueError("Donor ranking is tied")
    return {
        "q1": total_budget,
        "q2": jal_receipts,
        "q3": len(foreign_ids),
        "q4": ranking[0][0],
        "q5": round(total_receipts / total_budget * 100, 2),
    }


def read_sheet_xml(path: Path, sheet_name: str) -> dict[int, dict[int, object]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    doc_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet = next(
            node
            for node in workbook.findall(f".//{{{main_ns}}}sheet")
            if node.attrib["name"] == sheet_name
        )
        relation_id = sheet.attrib[f"{{{doc_rel_ns}}}id"]
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(
            node.attrib["Target"]
            for node in relationships.findall(f"{{{package_rel_ns}}}Relationship")
            if node.attrib["Id"] == relation_id
        )
        if target.startswith("/"):
            sheet_path = target.lstrip("/")
        elif target.startswith("xl/"):
            sheet_path = target
        else:
            sheet_path = posixpath.normpath(posixpath.join("xl", target))
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [
                "".join(text.text or "" for text in item.iter(f"{{{main_ns}}}t"))
                for item in shared_root.findall(f"{{{main_ns}}}si")
            ]
        root = ET.fromstring(archive.read(sheet_path))

    rows: dict[int, dict[int, object]] = {}
    for row_node in root.findall(f".//{{{main_ns}}}sheetData/{{{main_ns}}}row"):
        row_number = int(row_node.attrib["r"])
        values: dict[int, object] = {}
        for cell in row_node.findall(f"{{{main_ns}}}c"):
            reference = cell.attrib["r"]
            letters = "".join(char for char in reference if char.isalpha())
            column = 0
            for char in letters:
                column = column * 26 + ord(char.upper()) - ord("A") + 1
            cell_type = cell.attrib.get("t")
            value_node = cell.find(f"{{{main_ns}}}v")
            if cell_type == "inlineStr":
                inline = cell.find(f"{{{main_ns}}}is")
                value: object = (
                    "".join(text.text or "" for text in inline.iter(f"{{{main_ns}}}t"))
                    if inline is not None
                    else ""
                )
            elif cell_type == "s" and value_node is not None:
                value = shared[int(value_node.text)]
            elif value_node is None or value_node.text is None:
                value = None
            else:
                number = float(value_node.text)
                value = int(number) if number.is_integer() else number
            values[column] = value
        rows[row_number] = values
    return rows


def parse_and_derive_with_xml(path: Path) -> dict:
    rows = read_sheet_xml(path, "Finance FY25-26")
    budgets = [
        row
        for row in rows.values()
        if str(row.get(1, "")).startswith("BG-")
    ]
    receipts = [
        row
        for row in rows.values()
        if str(row.get(1, "")).startswith("RCPT-")
    ]
    if not budgets or not receipts:
        raise ValueError("OOXML parser found an empty detail table")
    donor_totals: dict[str, int] = defaultdict(int)
    for row in receipts:
        donor_totals[str(row[3])] += int(row[6])
    ranking = sorted(donor_totals.items(), key=lambda item: (-item[1], item[0]))
    if len(ranking) > 1 and ranking[0][1] == ranking[1][1]:
        raise ValueError("Donor ranking is tied")
    total_budget = sum(int(row[5]) for row in budgets)
    total_receipts = sum(int(row[6]) for row in receipts)
    return {
        "q1": total_budget,
        "q2": sum(
            int(row[6])
            for row in receipts
            if row[5] == "Jal Jeevan Saathi"
        ),
        "q3": len(
            {
                str(row[1])
                for row in receipts
                if row[4] == "Foreign Grant"
            }
        ),
        "q4": ranking[0][0],
        "q5": round(total_receipts / total_budget * 100, 2),
    }


def values_match(left, right, tol_pct: float = 0.0) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        return left == right
    allowed = max(1e-9, abs(float(right)) * tol_pct / 100)
    return math.isclose(float(left), float(right), rel_tol=0.0, abs_tol=allowed)


def main() -> int:
    try:
        case = json.loads(CASE_PATH.read_text(encoding="utf-8"))
        if case.get("id") != "two_tables":
            raise ValueError("case.json has the wrong case id")
        if len(case.get("questions", [])) not in (5, 6):
            raise ValueError("case.json must contain 5 or 6 questions")
        if len(case.get("files", [])) != 1:
            raise ValueError("case.json must name exactly one workbook")
        workbook = ROOT / "files" / case["files"][0]
        if not workbook.is_file() or workbook.stat().st_size >= 2 * 1024 * 1024:
            raise ValueError("Workbook is missing or is not under 2 MB")

        budgets, receipts = parse_with_openpyxl(workbook)
        first = derive_with_loops(budgets, receipts)
        second = parse_and_derive_with_xml(workbook)
        if set(first) != set(second):
            raise AssertionError("Independent derivations produced different question sets")
        for question_id in first:
            if not values_match(first[question_id], second[question_id]):
                raise AssertionError(
                    f"{question_id}: independent derivations disagree: "
                    f"{first[question_id]!r} != {second[question_id]!r}"
                )

        questions = case["questions"]
        expected_ids = {question["id"] for question in questions}
        if expected_ids != set(first) or len(expected_ids) != len(questions):
            raise AssertionError("Question IDs are missing, duplicated, or unexpected")
        for question in questions:
            actual = first[question["id"]]
            if not values_match(actual, question["answer"], float(question["tol_pct"])):
                raise AssertionError(
                    f"{question['id']}: expected {question['answer']!r}, got {actual!r}"
                )
        print(
            f"PASS two_tables: {len(questions)} questions; "
            "openpyxl and raw OOXML derivations agree"
        )
        return 0
    except Exception as exc:
        print(f"FAIL two_tables: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
