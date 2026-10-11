#!/usr/bin/env python3
"""Generate the deterministic two_tables benchmark case."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SEED = 20250417
ROOT = Path(__file__).resolve().parent
FILES = ROOT / "files"
WORKBOOK_NAME = "sahyog_budget_and_receipts.xlsx"


def make_clean_records() -> tuple[list[dict], list[dict]]:
    rng = np.random.default_rng(SEED)
    budget_dimensions = [
        ("Jal Jeevan Saathi", "Karnataka", "Raichur"),
        ("Jal Jeevan Saathi", "Karnataka", "Koppal"),
        ("Jal Jeevan Saathi", "Odisha", "Gajapati"),
        ("Jal Jeevan Saathi", "Odisha", "Koraput"),
        ("Poshan Sakhis", "Madhya Pradesh", "Mandla"),
        ("Poshan Sakhis", "Madhya Pradesh", "Dindori"),
        ("Poshan Sakhis", "Odisha", "Rayagada"),
        ("Digital Learning Centres", "Rajasthan", "Barmer"),
        ("Digital Learning Centres", "Rajasthan", "Jaisalmer"),
        ("Digital Learning Centres", "Karnataka", "Yadgir"),
        ("Climate-Resilient Livelihoods", "Maharashtra", "Nandurbar"),
        ("Climate-Resilient Livelihoods", "Maharashtra", "Gadchiroli"),
    ]
    budgets = []
    for index, (programme, state, district) in enumerate(budget_dimensions, start=1):
        approved = int(rng.integers(24, 72)) * 100_000
        community = int(rng.integers(2, 9)) * 25_000
        budgets.append(
            {
                "budget_code": f"BG-{index:03d}",
                "programme": programme,
                "state": state,
                "district": district,
                "approved_budget": approved,
                "community_contribution": community,
            }
        )

    donors = [
        ("Tata Trusts CSR", "CSR"),
        ("Azim Premji Philanthropic Initiatives", "Indian Foundation"),
        ("Jal Jeevan Mission", "Government Scheme"),
        ("Global Greengrants Fund", "Foreign Grant"),
        ("HDFC Bank Parivartan", "CSR"),
    ]
    programmes = [row["programme"] for row in budgets]
    receipts = []
    for index in range(1, 27):
        donor, donor_type = donors[(index - 1) % len(donors)]
        programme = programmes[(index * 5 + 1) % len(programmes)]
        amount = int(rng.integers(4, 18)) * 100_000 + int(rng.integers(0, 4)) * 25_000
        # Make the leading donor's aggregate unambiguously highest.
        if donor == "Tata Trusts CSR":
            amount += 900_000
        receipts.append(
            {
                "receipt_id": f"RCPT-25-{index:03d}",
                "date": datetime(2025, 4 + ((index - 1) % 6), 3 + ((index * 7) % 24)),
                "donor": donor,
                "donor_type": donor_type,
                "programme": programme,
                "amount": amount,
                "bank_reference": f"UTR{202500000 + index * 7919}",
            }
        )
    return budgets, receipts


def truth_from_clean(budgets: list[dict], receipts: list[dict]) -> dict:
    total_budget = sum(row["approved_budget"] for row in budgets)
    jal_receipts = sum(
        row["amount"] for row in receipts if row["programme"] == "Jal Jeevan Saathi"
    )
    foreign_count = len(
        {row["receipt_id"] for row in receipts if row["donor_type"] == "Foreign Grant"}
    )
    by_donor: dict[str, int] = {}
    for row in receipts:
        by_donor[row["donor"]] = by_donor.get(row["donor"], 0) + row["amount"]
    ranked = sorted(by_donor.items(), key=lambda item: (-item[1], item[0]))
    if len(ranked) > 1 and ranked[0][1] == ranked[1][1]:
        raise RuntimeError("Donor ranking unexpectedly tied")
    total_receipts = sum(row["amount"] for row in receipts)
    return {
        "q1": total_budget,
        "q2": jal_receipts,
        "q3": foreign_count,
        "q4": ranked[0][0],
        "q5": round(total_receipts / total_budget * 100, 2),
    }


def write_case_json(truth: dict) -> None:
    case = {
        "id": "two_tables",
        "description": (
            "A single messy finance sheet contains an approved-budget table and, "
            "after three blank rows, a donor-receipts table with a different schema."
        ),
        "files": [WORKBOOK_NAME],
        "questions": [
            {
                "id": "q1",
                "question": (
                    "What is the total approved FY 2025-26 budget in INR across all "
                    "programme-district rows? Do not add the displayed TOTAL APPROVED row."
                ),
                "answer": truth["q1"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "Adding the displayed total row double-counts the budget.",
            },
            {
                "id": "q2",
                "question": (
                    "How much donor receipt income in INR was credited for the "
                    "Jal Jeevan Saathi programme?"
                ),
                "answer": truth["q2"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "The receipts table has different columns from the budget table.",
            },
            {
                "id": "q3",
                "question": (
                    "How many distinct Receipt IDs in the donor-receipts table have "
                    "Donor Type exactly 'Foreign Grant'?"
                ),
                "answer": truth["q3"],
                "kind": "number",
                "tol_pct": 0.0,
                "trap": "Counting budget rows or the receipts total row gives a wrong count.",
            },
            {
                "id": "q4",
                "question": (
                    "Which donor has the highest total receipt amount when all of that "
                    "donor's receipt rows are combined?"
                ),
                "answer": truth["q4"],
                "kind": "name",
                "tol_pct": 0.0,
                "trap": "Selecting the largest single receipt is not the same as ranking donor totals.",
            },
            {
                "id": "q5",
                "question": (
                    "What percentage of the total approved budget has been received across "
                    "all donor receipts? Report 100 × total receipts / total approved budget, "
                    "rounded to two decimals, and exclude both displayed total rows."
                ),
                "answer": truth["q5"],
                "kind": "number",
                "tol_pct": 0.01,
                "trap": "Using either displayed total in addition to detail rows double-counts.",
            },
        ],
    }
    (ROOT / "case.json").write_text(
        json.dumps(case, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def style_header(ws, row: int, columns: int, fill: str) -> None:
    thin = Side(style="thin", color="A6A6A6")
    for cell in ws[row][:columns]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)


def write_workbook(budgets: list[dict], receipts: list[dict]) -> None:
    FILES.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Finance FY25-26"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    ws["A1"] = "SAHYOG RURAL INITIATIVE — PROGRAMME FINANCE REGISTER"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:G2")
    ws["A2"] = "FY 2025-26 | Internal management copy | Amounts in INR"
    ws["A2"].font = Font(italic=True, color="666666")

    ws.merge_cells("A4:F4")
    ws["A4"] = "A. APPROVED PROGRAMME BUDGET"
    ws["A4"].font = Font(bold=True, color="1F1F1F")
    ws["A4"].fill = PatternFill("solid", fgColor="D9EAF7")
    budget_headers = [
        "Budget Code",
        "Programme",
        "State",
        "District",
        "Approved Budget (INR)",
        "Community Contribution (INR)",
    ]
    ws.append(budget_headers)
    style_header(ws, 5, len(budget_headers), "4472C4")
    for record in budgets:
        ws.append(
            [
                record["budget_code"],
                record["programme"],
                record["state"],
                record["district"],
                record["approved_budget"],
                record["community_contribution"],
            ]
        )
    budget_total_row = ws.max_row + 1
    ws.cell(budget_total_row, 1, "TOTAL APPROVED")
    ws.cell(
        budget_total_row,
        5,
        f"=SUM(E6:E{budget_total_row - 1})",
    )
    ws.cell(
        budget_total_row,
        6,
        f"=SUM(F6:F{budget_total_row - 1})",
    )
    for cell in ws[budget_total_row][:6]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")

    # Exactly three blank rows separate the differently shaped tables.
    receipt_heading_row = budget_total_row + 4
    ws.merge_cells(
        start_row=receipt_heading_row,
        start_column=1,
        end_row=receipt_heading_row,
        end_column=7,
    )
    ws.cell(receipt_heading_row, 1, "B. DONOR RECEIPTS — CREDITS CLEARED APR–SEP 2025")
    ws.cell(receipt_heading_row, 1).font = Font(bold=True)
    ws.cell(receipt_heading_row, 1).fill = PatternFill("solid", fgColor="FFF2CC")
    receipt_headers = [
        "Receipt ID",
        "Credit Date",
        "Donor / Scheme",
        "Donor Type",
        "Programme Tagged",
        "Amount Received (INR)",
        "Bank Reference",
    ]
    receipt_header_row = receipt_heading_row + 1
    for col, value in enumerate(receipt_headers, start=1):
        ws.cell(receipt_header_row, col, value)
    style_header(ws, receipt_header_row, len(receipt_headers), "BF9000")
    for record in receipts:
        ws.append(
            [
                record["receipt_id"],
                record["date"],
                record["donor"],
                record["donor_type"],
                record["programme"],
                record["amount"],
                record["bank_reference"],
            ]
        )
    receipt_total_row = ws.max_row + 1
    ws.cell(receipt_total_row, 1, "TOTAL RECEIPTS")
    ws.cell(
        receipt_total_row,
        6,
        f"=SUM(F{receipt_header_row + 1}:F{receipt_total_row - 1})",
    )
    for cell in ws[receipt_total_row][:7]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6")

    for row in range(6, budget_total_row + 1):
        for col in (5, 6):
            ws.cell(row, col).number_format = '₹#,##0;[Red]-₹#,##0'
    for row in range(receipt_header_row + 1, receipt_total_row + 1):
        ws.cell(row, 2).number_format = "dd-mmm-yyyy"
        ws.cell(row, 6).number_format = '₹#,##0;[Red]-₹#,##0'
    widths = [18, 34, 20, 20, 32, 24, 22]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A{receipt_header_row}:G{receipt_total_row - 1}"

    notes = wb.create_sheet("Read me - notes")
    notes.column_dimensions["A"].width = 25
    notes.column_dimensions["B"].width = 95
    notes["A1"] = "Finance team notes"
    notes["A1"].font = Font(size=14, bold=True)
    notes_rows = [
        ("Workbook scope", "Approved budgets and cleared bank receipts through 30-Sep-2025."),
        ("Layout warning", "The Finance sheet contains two separate tables with different columns."),
        ("Totals", "Green and peach rows are display totals; do not add them to detail rows."),
        ("Amounts", "All monetary values are Indian rupees (INR), not lakhs."),
        ("Prepared by", "Programme Finance Unit, Sahyog Rural Initiative"),
    ]
    for key, value in notes_rows:
        notes.append([key, value])
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    path = FILES / WORKBOOK_NAME
    wb.save(path)
    if path.stat().st_size >= 2 * 1024 * 1024:
        raise RuntimeError(f"{path.name} exceeds 2 MB")


def main() -> None:
    budgets, receipts = make_clean_records()
    truth = truth_from_clean(budgets, receipts)
    # Truth and all questions are fixed before the records are rendered messily.
    write_case_json(truth)
    write_workbook(budgets, receipts)
    print(f"generated two_tables: {len(budgets)} budgets, {len(receipts)} receipts, 5 questions")


if __name__ == "__main__":
    main()
